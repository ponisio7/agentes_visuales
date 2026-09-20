# tests/test_docx_e2e_sin_llm.py
"""End-to-end sin LLM: un plan produce un .docx real con texto e imagen.

Reproduce la forma del fallo original (documento con texto + imagen) sin
depender del LLM ni de ningún caso concreto del problema:

  * plan -> ``PlanBuilder`` (corrector) -> ``PlanValidator`` -> sandbox/File.
  * El .docx se abre con python-docx, tiene un párrafo > 200 caracteres y al
    menos una imagen raster insertada.
  * Regresión del "contenido vacío": un alias ``respuesta`` que el LLM dejó
    suelto se resuelve al valor real y ya no llega ``'{}'``.
"""
import logging
import zipfile

import pytest
from docx import Document
from PIL import Image

from core.agent import EstadoAgente
from core.executors.dispatcher import AgentExecutor
from core.executors.file_executor import FileExecutor
from core.problem_solver.builder import PlanBuilder
from core.problem_solver.code_corrector import PythonCodeCorrector
from core.problem_solver.models import ExecutionPlan, StepPlan
from core.problem_solver.validator import PlanValidator

TEXTO = (
    "En una ciudad pequena habia una panaderia que abria antes del amanecer. "
    "Cada manana el panadero encendia el horno y dejaba que el olor despertara "
    "a los vecinos. Nadie sabia de donde venia la receta, pero todos coincidian "
    "en que sabia a domingo. "
) * 3


def _builder_validator():
    log = logging.getLogger("test.docx_e2e")
    log.addHandler(logging.NullHandler())
    validador = PlanValidator(log)
    return PlanBuilder(log, PythonCodeCorrector(), validador), validador


def _ejecutar_plan(plan):
    """Ejecuta los agentes del plan como lo hace el scheduler."""
    contexto = {}
    ultimo = (False, "", {})
    for agente in plan.agentes_generados:
        agente.estado = EstadoAgente.EJECUTANDO
        exito, mensaje, resultado = AgentExecutor.ejecutar(agente, contexto)
        ultimo = (exito, mensaje, resultado)
        if not exito:
            agente.estado = EstadoAgente.ERROR
            break
        agente.resultado = resultado
        agente.estado = EstadoAgente.COMPLETADO
        contexto[agente.nombre] = resultado
        contexto[agente.id] = resultado
    return contexto, ultimo


@pytest.fixture
def cwd_temporal(tmp_path, monkeypatch):
    """El FileExecutor solo acepta rutas relativas: cwd aislado."""
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _plan_documento(ruta_imagen: str) -> ExecutionPlan:
    """Plan genérico: un paso produce texto+imágenes y otro escribe el .docx."""
    plan = ExecutionPlan(problema_original="documento con texto e imagen")
    pasos = [
        StepPlan(
            orden=1,
            nombre="GenerarContenido",
            tipo_agente="Python",
            dependencia_ids=[],
            configuracion={"codigo": (
                f"resultado = {{'texto': {TEXTO!r}, 'imagenes': [{ruta_imagen!r}]}}"
            )},
        ),
        StepPlan(
            orden=2,
            nombre="EscribirDocumento",
            tipo_agente="File",
            dependencia_ids=["GenerarContenido"],
            configuracion={"operacion": "escribir", "archivo_destino": "documento.docx"},
        ),
    ]
    plan.pasos = pasos
    return plan


def test_pipeline_produce_docx_con_texto_e_imagen(cwd_temporal):
    imagen = cwd_temporal / "ilustracion.png"
    Image.new("RGB", (80, 60), "blue").save(imagen, format="PNG")

    builder, validador = _builder_validator()
    plan = _plan_documento(str(imagen))
    plan.agentes_generados = builder.generar_agentes(plan)

    ok, errores = validador.validar_plan(plan)
    assert ok, errores

    _, (exito, mensaje, resultado) = _ejecutar_plan(plan)
    assert exito is True, mensaje

    ruta_docx = cwd_temporal / "documento.docx"
    assert ruta_docx.exists()
    assert resultado.get("imagenes_insertadas", 0) >= 1

    documento = Document(str(ruta_docx))
    parrafos_largos = [p.text for p in documento.paragraphs if len(p.text) > 200]
    assert parrafos_largos, "el .docx no tiene ningún párrafo > 200 caracteres"
    assert len(documento.inline_shapes) >= 1, "el .docx no contiene imagen"

    with zipfile.ZipFile(ruta_docx) as paquete:
        medios = [n for n in paquete.namelist() if n.startswith("word/media/")]
    assert any(n.lower().endswith(".png") for n in medios), medios


def test_alias_respuesta_se_resuelve_al_valor_real(cwd_temporal):
    """Regresión del bug: el LLM escribe ``texto = respuesta`` y llegaba '{}'."""
    plan = ExecutionPlan(problema_original="resolver un alias suelto")
    plan.pasos = [
        StepPlan(
            orden=1,
            nombre="GenerarCuento",
            tipo_agente="Python",
            dependencia_ids=[],
            configuracion={"codigo": f"resultado = {{'cuento': {TEXTO!r}}}"},
        ),
        StepPlan(
            orden=2,
            nombre="CrearDocumento",
            tipo_agente="Python",
            dependencia_ids=["GenerarCuento"],
            configuracion={"codigo": (
                "texto = respuesta\n"
                "if not texto or len(str(texto)) < 200:\n"
                "    raise ValueError('contenido vacio o demasiado corto')\n"
                "resultado = {'documento': str(texto)}"
            )},
        ),
    ]

    builder, validador = _builder_validator()
    plan.agentes_generados = builder.generar_agentes(plan)

    # El código que va a ejecutarse (y que valida el validador) es el corregido.
    codigo_ejecutado = plan.pasos[1].configuracion["codigo"]
    assert "dependencia(contexto, 'GenerarCuento')" in codigo_ejecutado
    assert "get('body'" not in codigo_ejecutado

    ok, errores = validador.validar_plan(plan)
    assert ok, errores

    _, (exito, mensaje, resultado) = _ejecutar_plan(plan)
    assert exito is True, mensaje
    assert len(resultado["documento"]) > 200
    assert resultado["documento"] != "{}"


def test_docx_avisa_si_la_imagen_local_no_es_raster(cwd_temporal):
    """Un .png con SVG dentro no revienta: se reporta como imagen fallida."""
    from core.agent import Agente, TipoAgente

    falso_png = cwd_temporal / "dibujo.png"
    falso_png.write_text('<svg xmlns="http://www.w3.org/2000/svg"></svg>')

    agente = Agente(
        nombre="Escritor",
        tipo=TipoAgente.FILE,
        operacion_file="escribir",
        archivo_destino="documento.docx",
    )
    exito, mensaje, resultado = FileExecutor.ejecutar(
        agente,
        contexto={
            "GenerarContenido": {
                "texto": TEXTO,
                "imagenes": [str(falso_png)],
            }
        },
    )

    assert exito is True, mensaje
    assert resultado["imagenes_insertadas"] == 0
    assert resultado["imagenes_fallidas"], "debe reportar la imagen rechazada"
    assert "no contiene una imagen raster" in resultado["imagenes_fallidas"][0]
    assert (cwd_temporal / "documento.docx").exists()


def test_docx_inserta_imagen_aunque_la_clave_no_sea_imagenes(cwd_temporal):
    """El productor puede llamar a la clave 'imagen_ruta': se descubre por valor."""
    from core.agent import Agente, TipoAgente

    imagen = cwd_temporal / "ilustracion.png"
    Image.new("RGB", (40, 40), "green").save(imagen, format="PNG")

    agente = Agente(
        nombre="Escritor",
        tipo=TipoAgente.FILE,
        operacion_file="escribir",
        archivo_destino="documento.docx",
    )
    exito, mensaje, resultado = FileExecutor.ejecutar(
        agente,
        contexto={"GenerarContenido": {"texto": TEXTO, "imagen_ruta": str(imagen)}},
    )

    assert exito is True, mensaje
    assert resultado["imagenes_insertadas"] == 1
    documento = Document(str(cwd_temporal / "documento.docx"))
    assert len(documento.inline_shapes) == 1


def test_docx_no_confunde_un_texto_con_una_imagen(cwd_temporal):
    """Una descripción textual con 'imagen' en la clave no es una imagen."""
    from core.agent import Agente, TipoAgente

    agente = Agente(
        nombre="Escritor",
        tipo=TipoAgente.FILE,
        operacion_file="escribir",
        archivo_destino="documento.docx",
    )
    exito, mensaje, resultado = FileExecutor.ejecutar(
        agente,
        contexto={"GenerarContenido": {
            "texto": TEXTO,
            "descripcion_imagen": "un plato humeante sobre una mesa de madera",
        }},
    )

    assert exito is True, mensaje
    assert resultado["imagenes_insertadas"] == 0
    assert resultado["imagenes_fallidas"] == []


def test_validador_rechaza_pasos_file_imposibles(cwd_temporal, tmp_path):
    builder, validador = _builder_validator()

    plan = ExecutionPlan(problema_original="plan con pasos File imposibles")
    plan.pasos = [
        StepPlan(
            orden=1,
            nombre="Origen",
            tipo_agente="Python",
            dependencia_ids=[],
            configuracion={"codigo": "resultado = {'contenido': 'x'}"},
        ),
        StepPlan(
            orden=2,
            nombre="CopiaInutil",
            tipo_agente="File",
            dependencia_ids=["Origen"],
            configuracion={
                "operacion": "copiar",
                "archivo_origen": "documento.docx",
                "archivo_destino": "documento.docx",
            },
        ),
        StepPlan(
            orden=3,
            nombre="EscribeSinFuente",
            tipo_agente="File",
            dependencia_ids=[],
            configuracion={"operacion": "escribir", "archivo_destino": "otro.docx"},
        ),
    ]
    plan.agentes_generados = builder.generar_agentes(plan)
    ok, errores = validador.validar_plan(plan)

    assert ok is False
    assert any("SameFileError" in e for e in errores)
    assert any("sin dependencias" in e for e in errores)
