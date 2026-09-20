# tests/test_verification.py
"""Pruebas del verificador determinista de aceptación (H6).

Comprueba que la aceptación es una propiedad del ARTEFACTO (disco/bytes):
.docx sin imagen falla con motivo, .docx con imagen pasa, JSON no parseable
falla, imagen con extensión mentirosa falla, resultado vacío de un paso
crítico falla, y «sin contrato» se distingue de «verificado».
"""
import json

import pytest
from docx import Document
from PIL import Image

from core.agent import Agente, TipoAgente
from core.problem_solver.models import ContratoAceptacion
from core.verification import verificar_agente, verificar_contrato

TEXTO = (
    "Habia una vez una panaderia que abria antes del amanecer y dejaba que "
    "el olor del pan despertara a todo el barrio. "
) * 3


def _docx_con_imagen(ruta, png):
    documento = Document()
    documento.add_paragraph(TEXTO)
    documento.add_picture(str(png))
    documento.save(str(ruta))


def _docx_solo_texto(ruta):
    documento = Document()
    documento.add_paragraph(TEXTO)
    documento.save(str(ruta))


@pytest.fixture
def png_real(tmp_path):
    ruta = tmp_path / "ilustracion.png"
    Image.new("RGB", (60, 40), "green").save(ruta, format="PNG")
    return ruta


# ============================================================
# .docx CON / SIN IMAGEN
# ============================================================

def test_docx_sin_imagen_falla_con_motivo(tmp_path):
    ruta = tmp_path / "documento.docx"
    _docx_solo_texto(ruta)

    contrato = ContratoAceptacion(
        archivos=["documento.docx"], requiere_imagen=True, min_imagenes=1
    )
    resultado = verificar_contrato(contrato, {"texto": TEXTO}, cwd=str(tmp_path))

    assert resultado.aceptado is False
    assert resultado.verificado is True
    assert any("imágenes" in motivo or "imagenes" in motivo for motivo in resultado.motivos)
    assert "0 imágenes raster" in resultado.motivo()


def test_docx_con_imagen_pasa(tmp_path, png_real):
    ruta = tmp_path / "documento.docx"
    _docx_con_imagen(ruta, png_real)

    contrato = ContratoAceptacion(
        archivos=["documento.docx"], requiere_imagen=True, min_imagenes=1
    )
    resultado = verificar_contrato(contrato, {"texto": TEXTO}, cwd=str(tmp_path))

    assert resultado.aceptado is True, resultado.motivos
    assert resultado.verificado is True
    # La imagen incrustada se comprueba por bytes (PNG real).
    assert "PNG" in " ".join(c.detalle for c in resultado.comprobaciones)


def test_docx_sin_imagen_no_verificable_si_no_es_zip(tmp_path):
    (tmp_path / "documento.docx").write_text("esto no es un docx")

    contrato = ContratoAceptacion(
        archivos=["documento.docx"], requiere_imagen=True, min_imagenes=1
    )
    resultado = verificar_contrato(contrato, {}, cwd=str(tmp_path))

    assert resultado.aceptado is False
    assert resultado.motivos


# ============================================================
# JSON
# ============================================================

def test_json_no_parseable_falla(tmp_path):
    (tmp_path / "datos.json").write_text("{esto no es json", encoding="utf-8")

    contrato = ContratoAceptacion(archivos=["datos.json"], json_parseable=True)
    resultado = verificar_contrato(contrato, None, cwd=str(tmp_path))

    assert resultado.aceptado is False
    assert any("json" in motivo.lower() for motivo in resultado.motivos)


def test_json_parseable_pasa(tmp_path):
    (tmp_path / "datos.json").write_text(
        json.dumps({"a": 1, "b": 2}), encoding="utf-8"
    )

    contrato = ContratoAceptacion(
        archivos=["datos.json"], json_parseable=True, claves_requeridas=["a", "b"]
    )
    resultado = verificar_contrato(contrato, None, cwd=str(tmp_path))

    assert resultado.aceptado is True, resultado.motivos


def test_json_exigido_sin_fuente_no_se_da_por_bueno(tmp_path):
    contrato = ContratoAceptacion(archivos=["datos.txt"], json_parseable=True)
    (tmp_path / "datos.txt").write_text("hola", encoding="utf-8")

    resultado = verificar_contrato(contrato, None, cwd=str(tmp_path))

    assert resultado.aceptado is False
    assert any(c.no_verificable for c in resultado.comprobaciones)


# ============================================================
# IMÁGENES: FORMATO REAL, NO EXTENSIÓN
# ============================================================

def test_png_mentiroso_con_svg_dentro_falla(tmp_path):
    (tmp_path / "falso.png").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"></svg>', encoding="utf-8"
    )

    contrato = ContratoAceptacion(imagenes=["falso.png"])
    resultado = verificar_contrato(contrato, {}, cwd=str(tmp_path))

    assert resultado.aceptado is False
    assert any("raster" in motivo for motivo in resultado.motivos)


def test_formato_imagen_exigido(tmp_path, png_real):
    contrato_ok = ContratoAceptacion(
        imagenes=["ilustracion.png"], formato_imagen="PNG"
    )
    assert verificar_contrato(contrato_ok, {}, cwd=str(tmp_path)).aceptado is True

    contrato_mal = ContratoAceptacion(
        imagenes=["ilustracion.png"], formato_imagen="JPEG"
    )
    resultado = verificar_contrato(contrato_mal, {}, cwd=str(tmp_path))
    assert resultado.aceptado is False
    assert any("formato real" in motivo for motivo in resultado.motivos)


# ============================================================
# ARCHIVOS, TEXTO, ITEMS Y ERRORES
# ============================================================

def test_archivo_inexistente_falla(tmp_path):
    contrato = ContratoAceptacion(archivos=["no_existe.txt"])
    resultado = verificar_contrato(contrato, {}, cwd=str(tmp_path))

    assert resultado.aceptado is False
    assert any("no existe" in motivo for motivo in resultado.motivos)


def test_archivo_vacio_falla(tmp_path):
    (tmp_path / "vacio.txt").write_text("", encoding="utf-8")
    contrato = ContratoAceptacion(archivos=["vacio.txt"])
    resultado = verificar_contrato(contrato, {}, cwd=str(tmp_path))

    assert resultado.aceptado is False
    assert any("vacío" in motivo or "vacio" in motivo for motivo in resultado.motivos)


def test_ruta_absoluta_rechazada(tmp_path):
    contrato = ContratoAceptacion(archivos=["/etc/passwd"])
    resultado = verificar_contrato(contrato, {}, cwd=str(tmp_path))

    assert resultado.aceptado is False
    assert any("ruta no válida" in motivo for motivo in resultado.motivos)


def test_min_caracteres(tmp_path):
    contrato = ContratoAceptacion(min_caracteres=30)
    assert verificar_contrato(contrato, {"texto": "x" * 50}).aceptado is True
    assert verificar_contrato(contrato, {"texto": "corto"}).aceptado is False


def test_min_items_y_max_errores():
    contrato = ContratoAceptacion(min_items=3, max_errores=0)

    ok = verificar_contrato(contrato, {"items": [1, 2, 3], "errores": 0})
    assert ok.aceptado is True, ok.motivos

    mal = verificar_contrato(contrato, {"items": [1], "errores": [{"e": "x"}]})
    assert mal.aceptado is False
    assert any("items" in motivo for motivo in mal.motivos)
    assert any("errores" in motivo for motivo in mal.motivos)


def test_contrato_vacio_no_verifica():
    resultado = verificar_contrato(ContratoAceptacion(), {"status": "ok"})
    assert resultado.aceptado is True
    assert resultado.verificado is False


# ============================================================
# AGENTE: NIVEL 1 (es_critico) Y NIVEL 2 (contrato)
# ============================================================

def test_agente_sin_contrato_ni_critico_no_se_verifica():
    agente = Agente(nombre="Normal", tipo=TipoAgente.PYTHON)
    assert verificar_agente(agente, {"status": "ok"}) is None


def test_agente_critico_con_resultado_vacio_falla():
    agente = Agente(nombre="Critico", tipo=TipoAgente.PYTHON, es_critico=True)
    resultado = verificar_agente(agente, {})

    assert resultado is not None
    assert resultado.aceptado is False
    assert any("vacío" in motivo or "vacio" in motivo for motivo in resultado.motivos)


def test_agente_critico_con_resultado_util_pasa():
    agente = Agente(nombre="Critico", tipo=TipoAgente.PYTHON, es_critico=True)
    resultado = verificar_agente(agente, {"status": "ok", "datos": [1, 2]})

    assert resultado is not None
    assert resultado.aceptado is True, resultado.motivos


def test_agente_file_usa_su_destino_para_imagenes(tmp_path, png_real):
    """El contrato exige imagen y no lista el documento: se usa archivo_destino."""
    _docx_con_imagen(tmp_path / "doc.docx", png_real)
    agente = Agente(
        nombre="Escritor",
        tipo=TipoAgente.FILE,
        es_critico=True,
        operacion_file="escribir",
        archivo_destino="doc.docx",
        contrato_aceptacion={
            "requiere_imagen": True,
            "min_imagenes": 1,
        },
    )
    resultado = verificar_agente(agente, {"imagenes_insertadas": 1}, cwd=str(tmp_path))

    assert resultado is not None
    assert resultado.aceptado is True, resultado.motivos


def test_resultado_de_verificacion_es_serializable(tmp_path):
    contrato = ContratoAceptacion(archivos=["x.txt"])
    (tmp_path / "x.txt").write_text("hola", encoding="utf-8")
    resultado = verificar_contrato(contrato, {}, cwd=str(tmp_path))

    datos = resultado.to_dict()
    assert json.loads(json.dumps(datos))["aceptado"] is True


# ============================================================
# H6 REFINADO: MOTOR, DIRECTORIOS Y CONTENEDORES
# ============================================================

def test_verification_engine_fachada(tmp_path, png_real):
    from core.verification import VerificationEngine, VerificationResult

    engine = VerificationEngine(cwd=str(tmp_path))
    (tmp_path / "salida.txt").write_text("contenido", encoding="utf-8")

    resultado = engine.verificar_contrato(
        ContratoAceptacion(archivos=["salida.txt"])
    )

    assert isinstance(resultado, VerificationResult)
    assert resultado.ok is True
    assert resultado.estado == "verificado"
    assert "archivo:salida.txt" in resultado.criterios_comprobados
    assert resultado.evidencias
    assert "contenedor" in VerificationEngine.criterios_disponibles()


def test_resultado_no_verificable_se_marca(tmp_path):
    contrato = ContratoAceptacion(archivos=["x.pdf"], requiere_imagen=True)
    # .pdf con require_image pero sin validar_contenedor: se inspecciona.
    (tmp_path / "x.pdf").write_bytes(b"%PDF-1.4 no real")

    resultado = verificar_contrato(contrato, {}, cwd=str(tmp_path))

    assert resultado.aceptado is False
    assert resultado.estado == "fallido"
    assert resultado.criterios_fallidos


def test_directorio_con_archivos_esperados(tmp_path):
    carpeta = tmp_path / "salidas"
    carpeta.mkdir()
    (carpeta / "a.txt").write_text("a", encoding="utf-8")
    (carpeta / "b.txt").write_text("b", encoding="utf-8")

    ok = verificar_contrato(
        ContratoAceptacion(
            directorio="salidas", min_archivos=2, archivos_esperados=["a.txt", "b.txt"]
        ),
        {}, cwd=str(tmp_path),
    )
    assert ok.aceptado is True, ok.motivos

    falta = verificar_contrato(
        ContratoAceptacion(directorio="salidas", archivos_esperados=["c.txt"]),
        {}, cwd=str(tmp_path),
    )
    assert falta.aceptado is False
    assert any("faltan" in m for m in falta.motivos)


def test_directorio_inexistente_falla(tmp_path):
    resultado = verificar_contrato(
        ContratoAceptacion(directorio="no_existe", min_archivos=1),
        {}, cwd=str(tmp_path),
    )
    assert resultado.aceptado is False
    assert any("no existe el directorio" in m for m in resultado.motivos)


def test_validar_contenedor_docx(tmp_path, png_real):
    ruta = tmp_path / "documento.docx"
    _docx_con_imagen(ruta, png_real)

    ok = verificar_contrato(
        ContratoAceptacion(archivos=["documento.docx"], validar_contenedor=True),
        {}, cwd=str(tmp_path),
    )
    assert ok.aceptado is True, ok.motivos
    assert any("contenedor" in e for e in ok.evidencias)


def test_validar_contenedor_docx_corrupto(tmp_path):
    (tmp_path / "documento.docx").write_text("no soy un zip", encoding="utf-8")

    resultado = verificar_contrato(
        ContratoAceptacion(archivos=["documento.docx"], validar_contenedor=True),
        {}, cwd=str(tmp_path),
    )
    assert resultado.aceptado is False
    assert any("ZIP" in m or "contenedor" in m for m in resultado.motivos)


def test_contrato_from_dict_lee_campos_nuevos():
    contrato = ContratoAceptacion.from_dict({
        "directorio": "out",
        "min_archivos": "3",
        "archivos_esperados": ["a.txt"],
        "validar_contenedor": True,
    })
    assert contrato.directorio == "out"
    assert contrato.min_archivos == 3
    assert contrato.archivos_esperados == ["a.txt"]
    assert contrato.validar_contenedor is True
