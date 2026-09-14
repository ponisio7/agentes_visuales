# core/executors.py
"""
Ejecutores de agentes para el sistema Agentes Visuales.

Tipos soportados:
    - Python: ejecución en sandbox aislado
    - Shell:  comandos con detección automática de privilegios (pkexec)
    - HTTP:   peticiones con caché LRU, rate limiting y cancelación
    - File:   lectura/escritura con dispatch por extensión
              (.txt/.html/.csv/.md/.docx/.xlsx/.pdf) + copiar/mover/eliminar
    - LLM:    consultas a DeepSeek con parseo robusto de JSON
    - Loop:   procesamiento por lotes sobre listas

Características transversales:
    - Cancelación cooperativa vía CancellationToken
    - Caché LRU con TTL para HTTP
    - Rate limiting global
    - Progreso en tiempo real vía señales del bridge
    - Logging estructurado y errores accionables
"""

from __future__ import annotations

# ── Bibliotecas estándar ──
import atexit
import hashlib
import json
import logging
import os
import platform
import re
import shlex
import shutil
import subprocess
import threading
import time
from collections import OrderedDict
from typing import Any, Dict, List, Optional, Set, Tuple, Union
from urllib.parse import urlparse

# ── Terceros ──
import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

# ── Internos ──
from core.agent import Agente, TipoAgente
from core.sandbox import PythonSandbox, SandboxError
from .cancellation import CancellationToken


# ============================================================
# CONFIGURACIÓN DE LOGGING
# ============================================================
logger = logging.getLogger(__name__)


# ============================================================
# CONSTANTES
# ============================================================

# Límites de seguridad
MAX_BYTES_LECTURA_ARCHIVO = 10 * 1024 * 1024      # 10 MB
MAX_HTTP_BODY_SIZE = 50 * 1024 * 1024             # 50 MB
MAX_SHELL_COMMAND_LENGTH = 10000                  # 10 000 caracteres
MAX_FILE_PATH_LENGTH = 1000                       # 1 000 caracteres
MAX_CODIGO_LENGTH = 100000                        # 100 KB

# Caché HTTP
HTTP_CACHE_SIZE = 100
HTTP_CACHE_TTL = 300                              # 5 minutos
DEFAULT_USER_AGENT = "Agentes-Visuales/1.0"

# Métodos HTTP soportados
ALLOWED_HTTP_METHODS = frozenset(
    {"GET", "POST", "PUT", "DELETE", "PATCH", "HEAD", "OPTIONS"}
)

# Comandos shell que requieren privilegios de root
COMANDOS_PRIVILEGIADOS = frozenset({
    # Gestores de paquetes
    "apt", "apt-get", "aptitude", "dpkg", "snap", "yum", "dnf",
    "pacman", "zypper", "apk", "rpm",
    # Servicios y sistema
    "systemctl", "service", "init", "shutdown", "reboot", "halt",
    # Montaje y discos
    "mount", "umount", "fdisk", "parted", "mkfs", "fsck",
    "cryptsetup", "losetup",
    # Redes y firewall
    "iptables", "ip6tables", "nft", "ufw", "firewall-cmd",
    # Usuarios y permisos
    "useradd", "userdel", "usermod", "groupadd", "groupdel",
    "passwd", "chpasswd", "chsh", "visudo",
    # Modificación de sistema
    "modprobe", "insmod", "rmmod", "sysctl",
    # Otros
    "docker", "podman",
})

# Comandos shell peligrosos (solo warning, no bloqueo)
DANGEROUS_SHELL_COMMANDS = frozenset({
    "rm -rf", "dd if=", "mkfs", "chmod 777", "chown", "sudo",
    ":(){ :|:& };:", ">/dev/sda",
})

# Directorios peligrosos para eliminar
DANGEROUS_DIRS = frozenset({"/", "C:\\", "/home", "/root", "/etc", "/var"})

# Operaciones de File que escriben (para validación y normalización)
OPERACIONES_ESCRIBIR = frozenset({
    "escribir", "escribir_markdown", "escribir_docx",
    "escribir_xlsx", "escribir_pdf",
})

# Extensiones soportadas con su formato real
EXTENSIONES_ESCRITURA = {
    ".txt":  "texto",
    ".html": "texto",
    ".csv":  "texto",
    ".md":   "markdown",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".pdf":  "pdf",
}

# Claves prioritarias para extracción de contenido (fuente única de verdad)
_CLAVES_CONTENIDO_PRIORITARIAS = (
    "html",              # generación de páginas web
    "markdown",          # generación de documentos md
    "json",              # JSON parseado (LLM + HTTP)
    "contenido",         # File / Python explícito
    "respuesta_limpia",  # texto limpio sin fences
    "respuesta",         # LLM crudo
    "body",              # HTTP (crudo)
    "stdout",            # Shell
    "items",             # Loop
    "resultado",         # Python genérico
    "data",              # Genérico
    "texto",             # Genérico
    "output",            # Genérico
    "_av_json",          # fallback reservado
    "_av_respuesta_limpia",
    "_av_respuesta",
)

# Profundidad máxima de recursión al extraer contenido
_MAX_PROFUNDIDAD_EXTRACCION = 5


# ============================================================
# HELPERS GLOBALES
# ============================================================

def _limpiar_fences_markdown(texto: str) -> str:
    """Elimina ```json ... ``` y espacios."""
    if not texto:
        return ""
    t = texto.strip()
    t = re.sub(r"^```[a-zA-Z]*\s*", "", t)
    t = re.sub(r"\s*```\s*$", "", t)
    return t.strip()


def _parsear_json_robusto(texto: str) -> Optional[Any]:
    """Intenta parsear JSON con múltiples estrategias."""
    if not texto:
        return None
    try:
        return json.loads(texto)
    except Exception:
        pass
    try:
        ini, fin = texto.find("{"), texto.rfind("}")
        if ini != -1 and fin > ini:
            return json.loads(texto[ini:fin + 1])
    except Exception:
        pass
    try:
        ini, fin = texto.find("["), texto.rfind("]")
        if ini != -1 and fin > ini:
            return json.loads(texto[ini:fin + 1])
    except Exception:
        pass
    try:
        m = re.search(r"\{.*\}", texto, re.DOTALL)
        if m:
            return json.loads(m.group(0))
    except Exception:
        pass
    return None


def _es_resultado_sospechoso(resultado: Any) -> Tuple[bool, str]:
    """Detecta resultados vacíos que antes pasaban silenciosos."""
    if resultado is None:
        return True, "resultado es None"
    if resultado == {} or resultado == [] or resultado == "":
        return True, "resultado vacío"
    if isinstance(resultado, dict):
        vacios = (None, "", [], {}, "N/A", "null")
        if all(v in vacios for v in resultado.values()):
            return True, f"todos los valores vacíos: {list(resultado.keys())}"
        for k, v in resultado.items():
            if isinstance(v, dict) and not v:
                return True, f"clave '{k}' es dict vacío"
    return False, ""


def _sanitizar_para_reporte(valor: Any, max_len: int = 500) -> str:
    """Serializa un valor para mostrarlo en logs/errores, sin reventar."""
    try:
        if isinstance(valor, str):
            return valor[:max_len]
        return json.dumps(valor, default=str, ensure_ascii=False)[:max_len]
    except Exception:
        return str(valor)[:max_len]


def _celda_segura(valor: Any) -> Any:
    """Convierte un valor a algo que openpyxl acepte (str/int/float/bool/None)."""
    if valor is None:
        return ""
    if isinstance(valor, (dict, list)):
        return json.dumps(valor, ensure_ascii=False, default=str)
    if isinstance(valor, (int, float, bool, str)):
        return valor
    return str(valor)


def _escapar_xml(texto: str) -> str:
    """Escapa &, <, > para reportlab.Paragraph."""
    return (
        texto.replace("&", "&amp;")
             .replace("<", "&lt;")
             .replace(">", "&gt;")
    )

def _parece_csv(texto: str) -> bool:
    """Heurística: 2+ líneas con comas, sin HTML raro."""
    if not isinstance(texto, str):
        return False
    if "<br>" in texto:
        # A veces el LLM genera HTML con <br>. Los tratamos como saltos.
        texto = texto.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    lineas = [l for l in texto.strip().split("\n") if l.strip()]
    if len(lineas) < 2:
        return False
    lineas_con_coma = sum(1 for l in lineas if "," in l)
    return lineas_con_coma >= 2


def _parsear_csv_simple(texto: str) -> list:
    """Parsea un CSV simple a lista de listas. Sin quoting complejo."""
    import csv
    import io
    if "<br>" in texto:
        texto = texto.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    try:
        reader = csv.reader(io.StringIO(texto))
        filas = [row for row in reader if any(cell.strip() for cell in row)]
        # Asegurar que todas las filas tienen el mismo número de columnas
        if not filas:
            return []
        max_cols = max(len(f) for f in filas)
        filas_normalizadas = [
            f + [""] * (max_cols - len(f)) for f in filas
        ]
        return filas_normalizadas
    except Exception as e:
        logger.warning(f"Error parseando CSV: {e}")
        return []

# ============================================================
# CACHÉ HTTP (LRU con TTL)
# ============================================================

class HTTPCache:
    """Caché LRU thread-safe para resultados de peticiones HTTP."""

    def __init__(self, max_size: int = HTTP_CACHE_SIZE, ttl: int = HTTP_CACHE_TTL):
        self.max_size = max_size
        self.ttl = ttl
        self._cache: OrderedDict = OrderedDict()
        self._lock = threading.RLock()
        self._stats = {"hits": 0, "misses": 0}

    def _generate_key(self, url: str, method: str, headers: Dict, body: Optional[Any]) -> str:
        headers_norm = {k.lower(): v for k, v in (headers or {}).items()}
        headers_json = json.dumps(headers_norm, sort_keys=True)
        body_str = ""
        if body is not None:
            try:
                body_str = json.dumps(body, sort_keys=True, default=str)
            except (TypeError, ValueError):
                body_str = str(body)
        raw = f"{method.upper()}|{url}|{headers_json}|{body_str}"
        return hashlib.sha256(raw.encode("utf-8")).hexdigest()

    def get(self, url: str, method: str, headers: Dict, body: Optional[Any]) -> Optional[Dict]:
        key = self._generate_key(url, method, headers, body)
        with self._lock:
            if key not in self._cache:
                self._stats["misses"] += 1
                return None
            entry = self._cache[key]
            if time.time() - entry["timestamp"] > self.ttl:
                del self._cache[key]
                self._stats["misses"] += 1
                return None
            self._cache.move_to_end(key)
            self._stats["hits"] += 1
            return entry["result"]

    def put(self, url: str, method: str, headers: Dict, body: Optional[Any], result: Dict) -> None:
        key = self._generate_key(url, method, headers, body)
        with self._lock:
            if len(self._cache) >= self.max_size:
                oldest = next(iter(self._cache))
                del self._cache[oldest]
            self._cache[key] = {"timestamp": time.time(), "result": result}

    def clear(self) -> None:
        with self._lock:
            self._cache.clear()
            self._stats = {"hits": 0, "misses": 0}

    def get_stats(self) -> Dict[str, Union[int, float]]:
        with self._lock:
            total = self._stats["hits"] + self._stats["misses"]
            return {
                "size": len(self._cache),
                "hits": self._stats["hits"],
                "misses": self._stats["misses"],
                "hit_ratio": self._stats["hits"] / total if total > 0 else 0,
            }


# ============================================================
# RATE LIMITER
# ============================================================

class RateLimiter:
    """Rate limiter thread-safe (calls/segundo)."""

    def __init__(self, calls_per_second: float = 10):
        self.calls_per_second = calls_per_second
        self._last_call = 0.0
        self._lock = threading.RLock()

    def wait(self) -> None:
        with self._lock:
            now = time.time()
            elapsed = now - self._last_call
            min_interval = 1.0 / self.calls_per_second
            if elapsed < min_interval:
                time.sleep(min_interval - elapsed)
            self._last_call = time.time()

# ============================================================
# EJECUTOR PRINCIPAL
# ============================================================

class AgentExecutor:
    """Ejecuta agentes de diferentes tipos con funcionalidad real."""

    # ── Singletons ──
    _http_cache: Optional[HTTPCache] = None
    _rate_limiter: Optional[RateLimiter] = None
    _class_lock = threading.RLock()

    # ============================================================
    # RECURSOS COMPARTIDOS
    # ============================================================

    @classmethod
    def _get_http_cache(cls) -> HTTPCache:
        with cls._class_lock:
            if cls._http_cache is None:
                cls._http_cache = HTTPCache()
            return cls._http_cache

    @classmethod
    def _get_rate_limiter(cls) -> RateLimiter:
        with cls._class_lock:
            if cls._rate_limiter is None:
                cls._rate_limiter = RateLimiter()
            return cls._rate_limiter

    @classmethod
    def clear_http_cache(cls) -> None:
        with cls._class_lock:
            if cls._http_cache is not None:
                cls._http_cache.clear()

    @classmethod
    def get_http_cache_stats(cls) -> Dict:
        with cls._class_lock:
            if cls._http_cache is not None:
                return cls._http_cache.get_stats()
            return {"size": 0, "hits": 0, "misses": 0, "hit_ratio": 0}

    # ============================================================
    # PROGRESO
    # ============================================================

    @staticmethod
    def _actualizar_progreso(agente: Agente, progreso: int, mensaje: str = "") -> None:
        agente.progreso = min(100, max(0, progreso))
        if mensaje:
            agente.mensaje = mensaje
        if hasattr(agente, "_bridge") and agente._bridge is not None:
            try:
                agente._bridge.agente_actualizado.emit(agente.id)
            except Exception as e:
                logger.debug(f"Error emitiendo señal de progreso: {e}")

    # ============================================================
    # EXTRACCIÓN RECURSIVA DE CONTENIDO
    # ============================================================

    @staticmethod
    def _extraer_contenido_relevante(
        valor: Any,
        _profundidad: int = 0,
        _visitados: Optional[Set[int]] = None,
    ) -> Any:
        """
        Extrae el contenido más relevante de un resultado de agente.

        - Si 'valor' no es dict → devolverlo tal cual.
        - Si tiene una clave prioritaria con valor útil → devolverla.
        - Si tiene dicts hijos → recursar.
        - Si nada útil → devolver el dict original.
        """
        if _profundidad > _MAX_PROFUNDIDAD_EXTRACCION:
            return valor
        if valor is None or isinstance(valor, (str, int, float, bool, list)):
            return valor
        if not isinstance(valor, dict):
            return str(valor)

        if _visitados is None:
            _visitados = set()
        vid = id(valor)
        if vid in _visitados:
            return valor
        _visitados.add(vid)

        # 1. Clave prioritaria
        for clave in _CLAVES_CONTENIDO_PRIORITARIAS:
            if clave not in valor:
                continue
            encontrado = valor[clave]
            if encontrado is None:
                continue
            if isinstance(encontrado, str) and not encontrado.strip():
                continue
            if isinstance(encontrado, (list, dict)) and not encontrado:
                continue
            if isinstance(encontrado, (str, int, float, bool, list)):
                return encontrado
            if isinstance(encontrado, dict):
                rec = AgentExecutor._extraer_contenido_relevante(
                    encontrado, _profundidad + 1, _visitados
                )
                if rec is not encontrado:
                    return rec

        # 2. Dicts hijos
        for sub in valor.values():
            if isinstance(sub, dict) and sub:
                rec = AgentExecutor._extraer_contenido_relevante(
                    sub, _profundidad + 1, _visitados
                )
                if rec is not sub:
                    return rec

        return valor

    # ============================================================
    # VARIABLES Y SUSTITUCIÓN
    # ============================================================

    @staticmethod
    def _variables_disponibles(agente: Agente, contexto: Dict) -> Dict[str, str]:
        contexto = contexto or {}
        ahora = time.localtime()
        variables: Dict[str, str] = {
            "contexto": json.dumps(contexto, indent=2, default=str) if contexto else "",
            "resultado": json.dumps(contexto, indent=2, default=str) if contexto else "",
            "nombre": agente.nombre,
            "descripcion": agente.descripcion,
            "id": agente.id,
            "fecha": time.strftime("%Y-%m-%d", ahora),
            "hora": time.strftime("%H:%M:%S", ahora),
        }
        for clave, valor in contexto.items():
            if clave in variables:
                continue
            extraido = AgentExecutor._extraer_contenido_relevante(valor)
            if isinstance(extraido, str):
                variables[clave] = extraido
            elif isinstance(extraido, (dict, list)):
                try:
                    variables[clave] = json.dumps(
                        extraido, indent=2, default=str, ensure_ascii=False
                    )
                except (TypeError, ValueError):
                    variables[clave] = str(extraido)
            else:
                variables[clave] = str(extraido)
        return variables

    @staticmethod
    def _sustituir_variables(texto: str, variables: Dict[str, str]) -> str:
        if not texto or "{" not in texto or not variables:
            return texto
        pattern = re.compile(
            r"\{\s*(" + "|".join(re.escape(k) for k in variables.keys()) + r")\s*\}"
        )
        return pattern.sub(lambda m: variables.get(m.group(1), m.group(0)), texto)

    # ============================================================
    # VALIDADORES
    # ============================================================

    @staticmethod
    def _validar_ruta_archivo(ruta: str) -> bool:
        if not ruta or len(ruta) > MAX_FILE_PATH_LENGTH:
            return False
        normalized = os.path.normpath(ruta)
        if normalized.startswith(("..", "/", "\\")):
            return False
        basename = os.path.basename(normalized)
        if basename.startswith(".") and basename not in (".", ".."):
            return False
        return True

    @staticmethod
    def _validar_url(url: str) -> bool:
        if not url or len(url) > 2000:
            return False
        try:
            parsed = urlparse(url)
            if parsed.scheme not in ("http", "https"):
                return False
            if not parsed.netloc:
                return False
            if any(c in url for c in ("\n", "\r", "\t")):
                return False
            return True
        except Exception:
            return False

    @staticmethod
    def _es_lista_valida(items: Any) -> bool:
        return isinstance(items, list)

    @staticmethod
    def _resolver_ruta_en_contexto(contexto: Dict, ruta: str) -> Any:
        if not ruta:
            return None
        partes = ruta.split(".")
        valor = contexto.get(partes[0])
        for parte in partes[1:]:
            if isinstance(valor, dict):
                valor = valor.get(parte)
            else:
                return None
        return valor

    # ============================================================
    # COMANDOS PRIVILEGIADOS
    # ============================================================

    @staticmethod
    def _extraer_primer_comando(comando: str) -> Optional[str]:
        if not comando or not comando.strip():
            return None
        try:
            tokens = shlex.split(comando)
        except ValueError:
            tokens = comando.split()

        PALABRAS_RESERVADAS = frozenset({
            "if", "then", "else", "elif", "fi", "while", "do", "done",
            "for", "in", "case", "esac", "function", "return",
        })
        for token in tokens:
            if not token:
                continue
            if token in ("&&", "||", "|", ";", "&", "(", ")", "{", "}"):
                continue
            if token.startswith("-"):
                continue
            if "=" in token and not token.startswith(("/", "./")):
                izq, _, _ = token.partition("=")
                if izq.isidentifier():
                    continue
            if token in PALABRAS_RESERVADAS:
                continue
            return os.path.basename(token)
        return None

    @staticmethod
    def _comando_requiere_root(comando: str) -> bool:
        if not comando:
            return False
        cs = comando.strip()
        if cs.startswith(("sudo ", "pkexec ", "doas ")):
            return False
        primer = AgentExecutor._extraer_primer_comando(cs)
        return bool(primer) and primer in COMANDOS_PRIVILEGIADOS

    # ============================================================
    # DISPATCH PRINCIPAL
    # ============================================================

    @staticmethod
    def ejecutar(
        agente: Agente,
        contexto: Optional[Dict] = None,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        contexto = contexto or {}
        start_time = time.time()

        if cancellation_token and cancellation_token.esta_cancelado():
            return False, "Cancelado antes de iniciar", {"error": "cancelled"}

        try:
            es_valido, mensaje_error = agente.validar_configuracion()
            if not es_valido:
                return False, f"Configuración inválida: {mensaje_error}", {
                    "error": mensaje_error
                }

            tipo = agente.tipo
            dispatch = {
                TipoAgente.PYTHON: AgentExecutor._ejecutar_python,
                TipoAgente.SHELL:  AgentExecutor._ejecutar_shell,
                TipoAgente.HTTP:   AgentExecutor._ejecutar_http,
                TipoAgente.FILE:   AgentExecutor._ejecutar_file,
                TipoAgente.LLM:    AgentExecutor._ejecutar_llm,
                TipoAgente.LOOP:   AgentExecutor._ejecutar_loop,
            }
            handler = dispatch.get(tipo)
            if handler is None:
                return False, f"Tipo de agente no soportado: {tipo}", {}
            return handler(agente, contexto, cancellation_token)

        except Exception as e:
            execution_time = time.time() - start_time
            logger.exception(f"Error en ejecución de {agente.nombre}")
            return False, f"Error crítico: {e}", {
                "error": str(e),
                "tipo": type(e).__name__,
                "duracion": execution_time,
            }

    # ============================================================
    # 1. PYTHON
    # ============================================================

    @staticmethod
    def _ejecutar_python(
        agente: Agente,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        if cancellation_token and cancellation_token.esta_cancelado():
            return False, "Cancelado antes de ejecutar", {"error": "cancelled"}

        AgentExecutor._actualizar_progreso(agente, 20, "Preparando entorno Python...")
        codigo = agente.codigo_python or ""

        if not codigo.strip():
            AgentExecutor._actualizar_progreso(agente, 50, "Simulando ejecución...")
            duracion = max(0.01, float(agente.duracion or 0.5))
            for _ in range(int(min(duracion, 5) * 10)):
                if cancellation_token and cancellation_token.esta_cancelado():
                    return False, "Cancelado durante simulación", {"error": "cancelled"}
                time.sleep(0.1)
            AgentExecutor._actualizar_progreso(agente, 100, "Simulación completada")
            return True, f"Simulado {duracion:.2f}s", {
                "simulado": True, "duracion": duracion, "nombre": agente.nombre
            }

        if len(codigo) > MAX_CODIGO_LENGTH:
            AgentExecutor._actualizar_progreso(agente, 100, "Código demasiado largo")
            return False, f"Código demasiado largo (máx {MAX_CODIGO_LENGTH // 1024}KB)", {
                "error": "code_too_long"
            }

        timeout = getattr(agente, "timeout_python", 30)
        AgentExecutor._actualizar_progreso(agente, 40, "Ejecutando código Python...")

        try:
            exito, mensaje, resultado = PythonSandbox.ejecutar(
                codigo, contexto, timeout=timeout, cancellation_token=cancellation_token
            )
            if exito:
                sospechoso, razon = _es_resultado_sospechoso(resultado)
                if sospechoso:
                    logger.warning(
                        f"[{agente.nombre}] Resultado sospechoso: {razon} "
                        f"-> {_sanitizar_para_reporte(resultado, 200)}"
                    )
                    mensaje += f" (⚠ sospechoso: {razon})"

            AgentExecutor._actualizar_progreso(
                agente, 100,
                "Código ejecutado correctamente" if exito else f"Error: {mensaje[:50]}",
            )
            return exito, mensaje, resultado

        except SandboxError as e:
            AgentExecutor._actualizar_progreso(agente, 100, f"Error en sandbox: {str(e)[:50]}")
            return False, f"Error en sandbox: {e}", {"error": str(e)}

    # ============================================================
    # 2. SHELL
    # ============================================================

    @staticmethod
    def _ejecutar_shell(
        agente: Agente,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        if cancellation_token and cancellation_token.esta_cancelado():
            return False, "Cancelado antes de ejecutar", {"error": "cancelled"}

        AgentExecutor._actualizar_progreso(agente, 20, "Preparando comando shell...")

        variables = AgentExecutor._variables_disponibles(agente, contexto)
        comando = AgentExecutor._sustituir_variables(agente.comando_shell, variables)
        working_dir = AgentExecutor._sustituir_variables(agente.working_dir, variables)

        if not comando:
            AgentExecutor._actualizar_progreso(agente, 100, "Comando vacío")
            return False, "No hay comando shell definido", {}

        if len(comando) > MAX_SHELL_COMMAND_LENGTH:
            AgentExecutor._actualizar_progreso(agente, 100, "Comando demasiado largo")
            return False, (
                f"Comando demasiado largo (máx {MAX_SHELL_COMMAND_LENGTH} caracteres)"
            ), {}

        for peligro in DANGEROUS_SHELL_COMMANDS:
            if peligro in comando:
                logger.warning(f"Comando contiene operación peligrosa: {peligro}")

        comando_original = comando
        comando_ejecutable = comando
        usar_pkexec = False

        if AgentExecutor._comando_requiere_root(comando):
            pkexec_path = shutil.which("pkexec")
            if pkexec_path:
                comando_ejecutable = f"{pkexec_path} /bin/sh -c {shlex.quote(comando)}"
                usar_pkexec = True
                logger.info(f"[{agente.nombre}] pkexec aplicado: {comando[:60]}...")
                AgentExecutor._actualizar_progreso(agente, 35, "🔒 Solicitando privilegios...")
            else:
                mensaje_error = (
                    f"🔒 El comando requiere root, pero 'pkexec' no está instalado.\n"
                    f"Comando: {comando[:120]}\n"
                    f"Opciones:\n"
                    f"  1. sudo apt install policykit-1  (Debian/Ubuntu)\n"
                    f"  2. sudo dnf install polkit       (Fedora)\n"
                    f"  3. Antepón 'sudo' manualmente al comando."
                )
                logger.warning(f"[{agente.nombre}] {mensaje_error}")
                AgentExecutor._actualizar_progreso(agente, 100, "❌ Falta pkexec")
                return False, mensaje_error, {
                    "error": "requires_root_no_pkexec",
                    "comando": comando[:200],
                    "requiere_root": True,
                    "pkexec_disponible": False,
                }

        AgentExecutor._actualizar_progreso(agente, 50, f"Ejecutando: {comando[:50]}...")

        try:
            cwd = None
            if working_dir and AgentExecutor._validar_ruta_archivo(working_dir):
                if os.path.isdir(working_dir):
                    cwd = working_dir

            env = os.environ.copy()
            safe_vars = ["PATH", "HOME", "USER", "LANG", "LC_ALL", "TMPDIR"]
            if platform.system() == "Windows":
                safe_vars.extend(["SYSTEMROOT", "TEMP", "APPDATA"])
            clean_env = {k: env[k] for k in safe_vars if k in env}
            clean_env["PYTHONUNBUFFERED"] = "1"

            timeout = getattr(agente, "timeout_shell", 30)
            timeout_efectivo = timeout + 30 if usar_pkexec else timeout

            proceso = subprocess.Popen(
                comando_ejecutable,
                shell=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                cwd=cwd,
                env=clean_env,
                encoding="utf-8",
                errors="replace",
            )

            def cancelar_proceso(_token):
                try:
                    if proceso.poll() is None:
                        proceso.terminate()
                        time.sleep(0.3)
                        if proceso.poll() is None:
                            proceso.kill()
                        logger.info(f"Proceso shell cancelado: {comando[:50]}...")
                except Exception as e:
                    logger.warning(f"Error cancelando proceso shell: {e}")

            if cancellation_token:
                cancellation_token.agregar_callback(cancelar_proceso)

            try:
                inicio = time.time()
                while True:
                    if cancellation_token and cancellation_token.esta_cancelado():
                        cancelar_proceso(cancellation_token)
                        return False, "Cancelado por usuario", {
                            "error": "cancelled", "comando": comando[:100]
                        }
                    if time.time() - inicio > timeout_efectivo:
                        cancelar_proceso(cancellation_token)
                        return False, f"Timeout ({timeout_efectivo}s)", {
                            "error": "timeout", "comando": comando[:100]
                        }
                    if proceso.poll() is not None:
                        break
                    time.sleep(0.05)

                stdout, stderr = proceso.communicate(timeout=1)
                codigo = proceso.returncode
            finally:
                if cancellation_token:
                    cancellation_token.eliminar_callback(cancelar_proceso)

            AgentExecutor._actualizar_progreso(agente, 90, "Procesando resultado...")

            stdout = (stdout or "").strip()
            stderr = (stderr or "").strip()
            if len(stdout) > 10000:
                stdout = stdout[:10000] + "\n... (truncado)"
            if len(stderr) > 10000:
                stderr = stderr[:10000] + "\n... (truncado)"

            AgentExecutor._actualizar_progreso(agente, 100, "Comando completado")

            resultado_dict = {
                "codigo": codigo,
                "returncode": codigo,
                "exit_code": codigo,
                "codigo_salida": codigo,
                "stdout": stdout,
                "stderr": stderr,
                "comando_original": comando_original,
                "usado_pkexec": usar_pkexec,
            }

            if codigo != 0:
                error_lower = (stdout + " " + stderr).lower()
                patrones_permisos = (
                    "permission denied", "permiso denegado",
                    "operation not permitted", "are you root",
                    "must be root", "access denied",
                    "no se pudo abrir el fichero de bloqueo",
                )
                if any(p in error_lower for p in patrones_permisos):
                    resultado_dict["error_permisos"] = True
                    resultado_dict["sugerencia"] = (
                        "El comando falló por permisos. Verifica que:\n"
                        "  1. 'pkexec' está bien instalado\n"
                        "  2. Tu usuario está en el grupo adecuado\n"
                        "  3. El comando no está bloqueado por Polkit"
                    )

            if codigo == 0:
                mensaje = f"Comando ejecutado correctamente (código: {codigo})"
                if usar_pkexec:
                    mensaje += " [con pkexec]"
                if stdout:
                    mensaje += f"\nSalida: {stdout[:200]}..."
                return True, mensaje, resultado_dict
            else:
                mensaje = f"Comando falló (código: {codigo})"
                if usar_pkexec:
                    mensaje += " [con pkexec]"
                if stderr:
                    mensaje += f"\nError: {stderr[:200]}..."
                elif stdout:
                    mensaje += f"\nSalida: {stdout[:200]}..."
                return False, mensaje, resultado_dict

        except subprocess.TimeoutExpired:
            AgentExecutor._actualizar_progreso(agente, 100, "Timeout")
            return False, f"Comando excedió el tiempo límite ({timeout_efectivo}s)", {}
        except subprocess.SubprocessError as e:
            AgentExecutor._actualizar_progreso(agente, 100, f"Error: {str(e)[:50]}")
            return False, f"Error en subproceso: {e}", {}
        except Exception as e:
            AgentExecutor._actualizar_progreso(agente, 100, f"Error: {str(e)[:50]}")
            return False, f"Error en comando shell: {e}", {}

    # ============================================================
    # 3. HTTP
    # ============================================================

    @staticmethod
    def _ejecutar_http(
        agente: Agente,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        resultado_vacio = {
            "error": "cancelled", "cancelled": True,
            "status_code": None, "headers": {}, "body": "",
            "json": None, "url": "", "elapsed": 0,
        }
        if cancellation_token and cancellation_token.esta_cancelado():
            return False, "Cancelado antes de ejecutar", resultado_vacio

        AgentExecutor._actualizar_progreso(agente, 10, "Preparando petición HTTP...")
        variables = AgentExecutor._variables_disponibles(agente, contexto)

        url = AgentExecutor._sustituir_variables(agente.url_http or "", variables)
        if not url or not url.strip():
            AgentExecutor._actualizar_progreso(agente, 100, "URL vacía")
            return False, "No hay URL definida", {
                "error": "empty_url", "status_code": None, "headers": {},
                "body": "", "json": None, "url": "", "elapsed": 0,
            }
        if not AgentExecutor._validar_url(url):
            AgentExecutor._actualizar_progreso(agente, 100, "URL inválida")
            return False, f"URL inválida: {url}", {
                "error": "invalid_url", "status_code": None, "headers": {},
                "body": "", "json": None, "url": url[:100], "elapsed": 0,
            }

        metodo = (agente.metodo_http or "GET").upper()
        if metodo not in ALLOWED_HTTP_METHODS:
            AgentExecutor._actualizar_progreso(agente, 100, f"Método no soportado: {metodo}")
            return False, f"Método HTTP no soportado: {metodo}", {
                "error": "unsupported_method", "status_code": None,
                "headers": {}, "body": "", "json": None,
                "url": url[:100], "elapsed": 0,
            }

        headers: Dict[str, str] = {}
        if agente.headers_http:
            for k, v in agente.headers_http.items():
                if k and v:
                    try:
                        headers[k] = AgentExecutor._sustituir_variables(str(v), variables)
                    except Exception as e:
                        logger.warning(f"Error sustituyendo header '{k}': {e}")
                        headers[k] = str(v)
        headers.setdefault("User-Agent", DEFAULT_USER_AGENT)
        headers.setdefault("Accept", "application/json, */*")

        body_texto = AgentExecutor._sustituir_variables(agente.body_http or "", variables)
        body = None
        content_type = headers.get("Content-Type", "")

        if body_texto and metodo in ("POST", "PUT", "PATCH"):
            if "application/json" in content_type.lower() or body_texto.strip().startswith(("{", "[")):
                try:
                    body = json.loads(body_texto)
                except json.JSONDecodeError:
                    body = body_texto
                    headers.setdefault("Content-Type", "text/plain")
            else:
                body = body_texto
                headers.setdefault("Content-Type", "text/plain")

        timeout = getattr(agente, "timeout_http", 30)
        if timeout < 1:
            timeout = 30
            logger.warning(f"Timeout HTTP inválido, usando 30s")

        AgentExecutor._actualizar_progreso(agente, 30, f"{metodo} {url[:60]}...")

        # Caché
        http_cache = AgentExecutor._get_http_cache()
        cached = http_cache.get(url, metodo, headers, body)
        if cached is not None:
            AgentExecutor._actualizar_progreso(agente, 100, "✅ Respuesta desde caché")
            return True, "Respuesta desde caché", cached

        AgentExecutor._get_rate_limiter().wait()
        AgentExecutor._actualizar_progreso(agente, 50, "Enviando petición...")

        session = None
        response = None
        error = None
        completed = threading.Event()
        cancelar_peticion = None

        try:
            session = requests.Session()
            retry = Retry(
                total=2, backoff_factor=0.5,
                status_forcelist=[429, 500, 502, 503, 504],
                allowed_methods=list(ALLOWED_HTTP_METHODS),
            )
            session.mount("http://", HTTPAdapter(max_retries=retry))
            session.mount("https://", HTTPAdapter(max_retries=retry))

            def cancelar_peticion(_token):
                try:
                    if session:
                        session.close()
                        logger.info(f"🛑 Petición HTTP cancelada: {url[:60]}...")
                except Exception as e:
                    logger.warning(f"Error cancelando petición HTTP: {e}")

            if cancellation_token:
                cancellation_token.agregar_callback(cancelar_peticion)

            def hacer_peticion():
                nonlocal response, error
                try:
                    kwargs = {
                        "headers": headers,
                        "timeout": timeout,
                        "allow_redirects": True,
                        "verify": True,
                    }
                    if metodo in ("POST", "PUT", "PATCH"):
                        if isinstance(body, dict):
                            kwargs["json"] = body
                        else:
                            kwargs["data"] = body
                    method_fn = getattr(session, metodo.lower(), None)
                    if method_fn is None:
                        raise requests.exceptions.RequestException(
                            f"Método no soportado por requests: {metodo}"
                        )
                    response = method_fn(url, **kwargs)
                except Exception as e:
                    error = e
                finally:
                    completed.set()

            thread = threading.Thread(target=hacer_peticion, daemon=True)
            thread.start()

            inicio = time.time()
            while not completed.is_set():
                if cancellation_token and cancellation_token.esta_cancelado():
                    session.close()
                    AgentExecutor._actualizar_progreso(agente, 100, "⛔ Cancelado")
                    return False, "Cancelado por usuario", {
                        "error": "cancelled", "cancelled": True,
                        "status_code": None, "headers": {}, "body": "",
                        "json": None, "url": url[:100],
                        "elapsed": time.time() - inicio,
                    }
                if time.time() - inicio > timeout + 5:
                    session.close()
                    AgentExecutor._actualizar_progreso(agente, 100, "⏱ Timeout")
                    return False, f"Timeout en petición HTTP ({timeout}s)", {
                        "error": "timeout", "status_code": None, "headers": {},
                        "body": "", "json": None, "url": url[:100],
                        "elapsed": timeout,
                    }
                completed.wait(0.05)

            if error:
                raise error
            if response is None:
                raise requests.exceptions.RequestException("No se recibió respuesta")

        except requests.exceptions.Timeout:
            AgentExecutor._actualizar_progreso(agente, 100, "⏱ Timeout")
            return False, f"Timeout en petición HTTP ({timeout}s)", {
                "error": "timeout", "status_code": None, "headers": {},
                "body": "", "json": None, "url": url[:100], "elapsed": timeout,
            }
        except requests.exceptions.ConnectionError as e:
            AgentExecutor._actualizar_progreso(agente, 100, "❌ Error de conexión")
            return False, f"Error de conexión: {url}", {
                "error": "connection_error", "status_code": None,
                "headers": {}, "body": "", "json": None,
                "url": url[:100], "elapsed": 0, "detail": str(e),
            }
        except requests.exceptions.SSLError as e:
            AgentExecutor._actualizar_progreso(agente, 100, "❌ Error SSL")
            return False, f"Error SSL: {url}", {
                "error": "ssl_error", "status_code": None, "headers": {},
                "body": "", "json": None, "url": url[:100],
                "elapsed": 0, "detail": str(e),
            }
        except requests.exceptions.TooManyRedirects as e:
            AgentExecutor._actualizar_progreso(agente, 100, "❌ Demasiadas redirecciones")
            return False, f"Demasiadas redirecciones: {url}", {
                "error": "too_many_redirects", "status_code": None,
                "headers": {}, "body": "", "json": None,
                "url": url[:100], "elapsed": 0, "detail": str(e),
            }
        except requests.exceptions.RequestException as e:
            AgentExecutor._actualizar_progreso(agente, 100, f"❌ Error: {str(e)[:50]}")
            return False, f"Error en petición HTTP: {e}", {
                "error": "request_exception", "status_code": None,
                "headers": {}, "body": "", "json": None,
                "url": url[:100], "elapsed": 0, "detail": str(e),
            }
        except Exception as e:
            AgentExecutor._actualizar_progreso(agente, 100, f"❌ Error inesperado")
            logger.exception(f"Error inesperado en HTTP: {e}")
            return False, f"Error inesperado en HTTP: {e}", {
                "error": "unexpected", "status_code": None, "headers": {},
                "body": "", "json": None, "url": url[:100],
                "elapsed": 0, "detail": str(e),
            }
        finally:
            if cancellation_token and cancelar_peticion:
                cancellation_token.eliminar_callback(cancelar_peticion)

        AgentExecutor._actualizar_progreso(
            agente, 80, f"📥 Procesando respuesta {response.status_code}..."
        )

        body_size = len(response.content)
        body_truncated = body_size > MAX_HTTP_BODY_SIZE

        try:
            body_text = response.text
        except UnicodeDecodeError:
            import base64
            body_text = (
                f"[BINARY DATA: "
                f"{base64.b64encode(response.content).decode('ascii')[:100]}...]"
            )

        if body_truncated:
            body_text = body_text[:10000] + "\n... (truncado)"

        json_data = None
        if response.headers.get("Content-Type", "").startswith("application/json"):
            try:
                json_data = response.json()
            except (ValueError, json.JSONDecodeError):
                try:
                    json_data = json.loads(body_text[:10000])
                except (ValueError, json.JSONDecodeError):
                    pass

        resultado = {
            "status_code": response.status_code,
            "headers": dict(response.headers),
            "body": body_text,
            "body_truncated": body_truncated,
            "body_size": body_size,
            "json": json_data,
            "url": response.url,
            "elapsed": response.elapsed.total_seconds(),
            "error": None,
        }

        if 200 <= response.status_code < 400:
            http_cache.put(url, metodo, headers, body, resultado)

        AgentExecutor._actualizar_progreso(
            agente, 100, f"✅ Completado (status: {response.status_code})"
        )

        if 200 <= response.status_code < 300:
            mensaje = f"Petición exitosa (status: {response.status_code})"
            if isinstance(json_data, dict):
                claves = list(json_data.keys())[:3]
                mensaje += f" - JSON claves: {claves}"
            elif isinstance(json_data, list):
                mensaje += f" - JSON array: {len(json_data)} items"
            return True, mensaje, resultado
        else:
            mensaje = f"Petición falló (status: {response.status_code})"
            if isinstance(json_data, dict):
                err = (
                    json_data.get("error")
                    or json_data.get("message")
                    or json_data.get("detail")
                )
                if err:
                    mensaje += f" - {err[:100]}"
            return False, mensaje, resultado

    # ============================================================
    # 4. FILE — DISPATCH POR EXTENSIÓN
    # ============================================================

    @staticmethod
    def _ejecutar_file(
        agente: Agente,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        """
        Operaciones con archivos.

        Para escribir, acepta tanto operacion='escribir' como
        'escribir_docx'/'escribir_xlsx'/'escribir_pdf'/'escribir_markdown'.
        En todos los casos, el formato real lo decide la extensión de
        'archivo_destino' (dispatch por extensión).
        """
        if cancellation_token and cancellation_token.esta_cancelado():
            return False, "Cancelado antes de ejecutar", {"error": "cancelled"}

        AgentExecutor._actualizar_progreso(agente, 20, "Preparando operación de archivo...")

        variables = AgentExecutor._variables_disponibles(agente, contexto)
        origen = AgentExecutor._sustituir_variables(agente.archivo_origen, variables)
        destino = AgentExecutor._sustituir_variables(agente.archivo_destino, variables)
        operacion = agente.operacion_file or "leer"

        AgentExecutor._actualizar_progreso(agente, 40, f"Operación: {operacion}")

        # ── Validación según operación ──
        if operacion in ("leer", "eliminar"):
            ruta_archivo = origen
            if not ruta_archivo:
                AgentExecutor._actualizar_progreso(agente, 100, "Ruta no especificada")
                return False, f"No se especificó ruta para: {operacion}", {}
            if not AgentExecutor._validar_ruta_archivo(ruta_archivo):
                AgentExecutor._actualizar_progreso(agente, 100, "Ruta inválida")
                return False, f"Ruta inválida: {ruta_archivo}", {}

        elif operacion in OPERACIONES_ESCRIBIR or operacion in ("copiar", "mover"):
            if not destino:
                AgentExecutor._actualizar_progreso(agente, 100, "Destino no especificado")
                return False, f"No se especificó destino para: {operacion}", {}
            if not AgentExecutor._validar_ruta_archivo(destino):
                AgentExecutor._actualizar_progreso(agente, 100, "Ruta destino inválida")
                return False, f"Ruta destino inválida: {destino}", {}

            if operacion in ("copiar", "mover"):
                if not origen or not AgentExecutor._validar_ruta_archivo(origen):
                    AgentExecutor._actualizar_progreso(agente, 100, "Ruta origen inválida")
                    return False, f"Ruta origen inválida: {origen}", {}
                ruta_archivo = origen
            else:
                ruta_archivo = destino

            # ✅ Normalizar "escribir_X" → "escribir". El formato real
            # lo decide la extensión de 'archivo_destino' en el dispatch.
            if operacion != "escribir" and operacion in OPERACIONES_ESCRIBIR:
                logger.debug(
                    f"Normalizando operación '{operacion}' → 'escribir' "
                    f"(archivo_destino='{destino}')"
                )
                operacion = "escribir"
        else:
            AgentExecutor._actualizar_progreso(agente, 100, f"Operación no soportada: {operacion}")
            return False, f"Operación de archivo no soportada: {operacion}", {}

        if cancellation_token and cancellation_token.esta_cancelado():
            return False, "Cancelado antes de operación", {"error": "cancelled"}

        try:
            if operacion == "leer":
                return AgentExecutor._file_leer(agente, ruta_archivo, cancellation_token)

            if operacion == "escribir":
                contenido, error = AgentExecutor._extraer_contenido_para_escritura(
                    contexto, operacion
                )
                if error is not None:
                    _, msg_err, res_err = error
                    res_err["archivo"] = ruta_archivo
                    AgentExecutor._actualizar_progreso(agente, 100, "❌ Sin contenido")
                    return False, msg_err, res_err

                # Dispatch por extensión
                _, ext = os.path.splitext(ruta_archivo.lower())
                if ext == ".md":
                    return AgentExecutor._file_escribir_markdown(
                        agente, ruta_archivo, contenido, contexto, cancellation_token
                    )
                if ext == ".docx":
                    return AgentExecutor._file_escribir_docx(
                        agente, ruta_archivo, contenido, contexto, cancellation_token
                    )
                if ext == ".xlsx":
                    return AgentExecutor._file_escribir_xlsx(
                        agente, ruta_archivo, contenido, contexto, cancellation_token
                    )
                if ext == ".pdf":
                    return AgentExecutor._file_escribir_pdf(
                        agente, ruta_archivo, contenido, contexto, cancellation_token
                    )
                return AgentExecutor._file_escribir_texto(
                    agente, ruta_archivo, contenido, contexto, cancellation_token
                )

            if operacion == "copiar":
                return AgentExecutor._file_copiar(agente, origen, destino)

            if operacion == "mover":
                return AgentExecutor._file_mover(agente, origen, destino)

            if operacion == "eliminar":
                return AgentExecutor._file_eliminar(agente, ruta_archivo)

            AgentExecutor._actualizar_progreso(agente, 100, f"Operación no soportada: {operacion}")
            return False, f"Operación de archivo no soportada: {operacion}", {}

        except PermissionError as e:
            AgentExecutor._actualizar_progreso(agente, 100, "Permiso denegado")
            return False, (
                f"Permiso denegado para '{operacion}' en: {ruta_archivo} - {e}"
            ), {
                "error": "permission_denied",
                "archivo": ruta_archivo,
                "operacion": operacion,
            }
        except OSError as e:
            AgentExecutor._actualizar_progreso(agente, 100, f"Error: {str(e)[:50]}")
            return False, f"Error en operación de archivo: {e}", {
                "error": "os_error",
                "archivo": ruta_archivo,
                "operacion": operacion,
                "detalle": str(e),
            }
        except Exception as e:
            AgentExecutor._actualizar_progreso(agente, 100, f"Error: {str(e)[:50]}")
            return False, f"Error inesperado en operación de archivo: {e}", {
                "error": "unexpected",
                "archivo": ruta_archivo,
                "operacion": operacion,
                "detalle": str(e),
            }

    # ── Sub-operaciones de File ──

    @staticmethod
    def _file_leer(
        agente: Agente,
        ruta_archivo: str,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        if not os.path.exists(ruta_archivo):
            AgentExecutor._actualizar_progreso(agente, 100, "Archivo no encontrado")
            return False, f"Archivo no encontrado: {ruta_archivo}", {}
        if not os.path.isfile(ruta_archivo):
            AgentExecutor._actualizar_progreso(agente, 100, "No es un archivo")
            return False, f"No es un archivo: {ruta_archivo}", {}

        tamaño = os.path.getsize(ruta_archivo)
        if tamaño > MAX_BYTES_LECTURA_ARCHIVO:
            AgentExecutor._actualizar_progreso(agente, 100, "Archivo demasiado grande")
            return False, (
                f"Archivo demasiado grande ({tamaño} > {MAX_BYTES_LECTURA_ARCHIVO})"
            ), {"archivo": ruta_archivo, "tamaño": tamaño}

        AgentExecutor._actualizar_progreso(agente, 70, "Leyendo archivo...")
        contenido = ""
        with open(ruta_archivo, "r", encoding="utf-8", errors="replace") as f:
            while True:
                if cancellation_token and cancellation_token.esta_cancelado():
                    return False, "Cancelado durante lectura", {
                        "error": "cancelled",
                        "archivo": ruta_archivo,
                        "bytes_leidos": len(contenido),
                    }
                chunk = f.read(8192)
                if not chunk:
                    break
                contenido += chunk

        AgentExecutor._actualizar_progreso(agente, 100, "Archivo leído")

        resultado = {
            "archivo": ruta_archivo,
            "tamaño": tamaño,
            "contenido": contenido,
            "total_caracteres": len(contenido),
        }
        if ruta_archivo.endswith(".json"):
            try:
                resultado["json"] = json.loads(contenido)
            except json.JSONDecodeError:
                pass

        modo_salida = getattr(agente, "modo_salida_file", "auto") or "auto"
        resultado = AgentExecutor._aplicar_modo_salida_file(resultado, modo_salida, contenido)

        return True, f"Archivo leído: {ruta_archivo} ({len(contenido)} caracteres)", resultado

    @staticmethod
    def _file_escribir_texto(
        agente: Agente,
        ruta_archivo: str,
        contenido: Any,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        if isinstance(contenido, str):
            contenido_str = contenido
        elif isinstance(contenido, (dict, list)):
            contenido_str = json.dumps(contenido, indent=2, default=str, ensure_ascii=False)
        elif isinstance(contenido, (int, float, bool)):
            contenido_str = str(contenido)
        else:
            contenido_str = str(contenido)

        if not contenido_str.strip():
            return False, f"File.escribir_texto: contenido vacío para '{ruta_archivo}'", {
                "error": "empty_content", "archivo": ruta_archivo
            }

        directorio = os.path.dirname(ruta_archivo)
        if directorio:
            os.makedirs(directorio, exist_ok=True)

        AgentExecutor._actualizar_progreso(agente, 70, "Escribiendo archivo...")

        try:
            with open(ruta_archivo, "w", encoding="utf-8") as f:
                for i in range(0, len(contenido_str), 8192):
                    if cancellation_token and cancellation_token.esta_cancelado():
                        return False, "Cancelado durante escritura", {
                            "error": "cancelled",
                            "archivo": ruta_archivo,
                            "caracteres_escritos": i,
                        }
                    f.write(contenido_str[i:i + 8192])
        except OSError as e:
            return False, f"File.escribir_texto: error al escribir '{ruta_archivo}': {e}", {
                "error": "os_error", "archivo": ruta_archivo, "detalle": str(e)
            }

        if not os.path.exists(ruta_archivo) or os.path.getsize(ruta_archivo) == 0:
            return False, f"File.escribir_texto: '{ruta_archivo}' no se creó o está vacío", {
                "error": "write_failed", "archivo": ruta_archivo
            }

        tamaño_real = os.path.getsize(ruta_archivo)
        resultado = {
            "archivo": ruta_archivo,
            "ruta_absoluta": os.path.abspath(ruta_archivo),
            "caracteres_escritos": len(contenido_str),
            "tamaño": tamaño_real,
            "bytes_en_disco": tamaño_real,
            "contenido": contenido_str,
            "contenido_preview": contenido_str[:200] if len(contenido_str) > 200 else contenido_str,
        }
        modo_salida = getattr(agente, "modo_salida_file", "auto") or "auto"
        resultado = AgentExecutor._aplicar_modo_salida_file(resultado, modo_salida, contenido_str)

        AgentExecutor._actualizar_progreso(agente, 100, "Archivo escrito")
        return True, (
            f"Archivo escrito: {ruta_archivo} ({len(contenido_str)} caracteres, "
            f"{tamaño_real} bytes)"
        ), resultado

    @staticmethod
    def _file_escribir_markdown(
        agente: Agente,
        ruta_archivo: str,
        contenido: Any,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        base, ext = os.path.splitext(ruta_archivo)
        if not ext:
            ruta_archivo = base + ".md"
        return AgentExecutor._file_escribir_texto(
            agente, ruta_archivo, contenido, contexto, cancellation_token
        )

    @staticmethod
    def _file_escribir_docx(
        agente: Agente,
        ruta_archivo: str,
        contenido: Any,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return False, (
                "python-docx no está instalado. Instálalo con: "
                "pip install python-docx"
            ), {"error": "missing_dependency", "dep": "python-docx"}

        directorio = os.path.dirname(ruta_archivo)
        if directorio:
            os.makedirs(directorio, exist_ok=True)

        AgentExecutor._actualizar_progreso(agente, 70, "Escribiendo .docx...")

        doc = Document()

        titulo = None
        if isinstance(contenido, dict):
            titulo = contenido.get("titulo") or contenido.get("title")
            for clave in ("contenido", "texto", "respuesta_limpia", "respuesta"):
                if clave in contenido:
                    contenido = contenido[clave]
                    break

        if titulo:
            h = doc.add_heading(str(titulo), level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if isinstance(contenido, str):
            texto = contenido
        elif isinstance(contenido, (dict, list)):
            texto = json.dumps(contenido, indent=2, ensure_ascii=False, default=str)
        else:
            texto = str(contenido)

        if not texto.strip():
            return False, f"File.escribir_docx: contenido vacío para '{ruta_archivo}'", {
                "error": "empty_content", "archivo": ruta_archivo
            }

        for linea in texto.split("\n"):
            if cancellation_token and cancellation_token.esta_cancelado():
                return False, "Cancelado durante generación de .docx", {
                    "error": "cancelled", "archivo": ruta_archivo
                }
            stripped = linea.strip()
            if not stripped:
                doc.add_paragraph()
                continue
            if stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith(("- ", "* ")):
                doc.add_paragraph(stripped[2:], style="List Bullet")
            elif re.match(r"^\d+\.\s", stripped):
                doc.add_paragraph(re.sub(r"^\d+\.\s", "", stripped), style="List Number")
            else:
                doc.add_paragraph(stripped)

        try:
            doc.save(ruta_archivo)
        except OSError as e:
            return False, f"File.escribir_docx: error al guardar '{ruta_archivo}': {e}", {
                "error": "os_error", "archivo": ruta_archivo, "detalle": str(e)
            }

        if not os.path.exists(ruta_archivo) or os.path.getsize(ruta_archivo) == 0:
            return False, f"File.escribir_docx: '{ruta_archivo}' no se creó o está vacío", {
                "error": "write_failed", "archivo": ruta_archivo
            }

        tamaño = os.path.getsize(ruta_archivo)
        AgentExecutor._actualizar_progreso(agente, 100, "Archivo .docx escrito")

        return True, f"Archivo .docx escrito: {ruta_archivo} ({tamaño} bytes)", {
            "archivo": ruta_archivo,
            "ruta_absoluta": os.path.abspath(ruta_archivo),
            "tamaño": tamaño,
            "bytes_en_disco": tamaño,
            "contenido": texto[:500],
            "formato": "docx",
        }

    @staticmethod
    def _file_escribir_xlsx(
        agente: Agente,
        ruta_archivo: str,
        contenido: Any,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            return False, (
                "openpyxl no está instalado. Instálalo con: pip install openpyxl"
            ), {"error": "missing_dependency", "dep": "openpyxl"}

        directorio = os.path.dirname(ruta_archivo)
        if directorio:
            os.makedirs(directorio, exist_ok=True)

        AgentExecutor._actualizar_progreso(agente, 70, "Escribiendo .xlsx...")

        # ── ✅ NUEVO: autodetectar string CSV y parsearlo ──
        if isinstance(contenido, str) and _parece_csv(contenido):
            filas_csv = _parsear_csv_simple(contenido)
            if filas_csv:
                logger.info(
                    f"File.escribir_xlsx: contenido era string CSV, "
                    f"parseado a {len(filas_csv)} filas x "
                    f"{max(len(f) for f in filas_csv)} columnas"
                )
                contenido = filas_csv

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Datos"

        if isinstance(contenido, list) and contenido:
            if isinstance(contenido[0], dict):
                # Lista de dicts → una fila por dict, cabeceras = claves
                headers = list(contenido[0].keys())
                ws.append(headers)
                for col_idx, _ in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(
                        start_color="4472C4", end_color="4472C4",
                        fill_type="solid"
                    )
                    cell.alignment = Alignment(horizontal="center")
                for row in contenido:
                    if cancellation_token and cancellation_token.esta_cancelado():
                        return False, "Cancelado durante .xlsx", {
                            "error": "cancelled", "archivo": ruta_archivo
                        }
                    ws.append([_celda_segura(row.get(h, "")) for h in headers])
            elif isinstance(contenido[0], (list, tuple)):
                # Lista de listas → la primera fila son las cabeceras
                ws.append([_celda_segura(v) for v in contenido[0]])
                for col_idx, _ in enumerate(contenido[0], start=1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(
                        start_color="4472C4", end_color="4472C4",
                        fill_type="solid"
                    )
                    cell.alignment = Alignment(horizontal="center")
                for row in contenido[1:]:
                    if cancellation_token and cancellation_token.esta_cancelado():
                        return False, "Cancelado durante .xlsx", {
                            "error": "cancelled", "archivo": ruta_archivo
                        }
                    ws.append([_celda_segura(v) for v in row])
            else:
                # Lista plana de strings/números → una columna
                for item in contenido:
                    ws.append([_celda_segura(item)])
        elif isinstance(contenido, dict):
            # Dict → clave/valor en dos columnas
            for k, v in contenido.items():
                ws.append([str(k), _celda_segura(v)])
        else:
            # String plano sin comas / número / None → una celda
            ws.append([_celda_segura(contenido)])

        try:
            wb.save(ruta_archivo)
        except OSError as e:
            return False, f"File.escribir_xlsx: error al guardar '{ruta_archivo}': {e}", {
                "error": "os_error", "archivo": ruta_archivo, "detalle": str(e)
            }

        if not os.path.exists(ruta_archivo) or os.path.getsize(ruta_archivo) == 0:
            return False, f"File.escribir_xlsx: '{ruta_archivo}' no se creó o está vacío", {
                "error": "write_failed", "archivo": ruta_archivo
            }

        tamaño = os.path.getsize(ruta_archivo)
        AgentExecutor._actualizar_progreso(agente, 100, "Archivo .xlsx escrito")

        return True, f"Archivo .xlsx escrito: {ruta_archivo} ({tamaño} bytes)", {
            "archivo": ruta_archivo,
            "ruta_absoluta": os.path.abspath(ruta_archivo),
            "tamaño": tamaño,
            "bytes_en_disco": tamaño,
            "formato": "xlsx",
        }
    
    
    @staticmethod
    def _file_escribir_pdf(
        agente: Agente,
        ruta_archivo: str,
        contenido: Any,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        try:
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import cm
        except ImportError:
            return False, (
                "reportlab no está instalado. Instálalo con: pip install reportlab"
            ), {"error": "missing_dependency", "dep": "reportlab"}

        directorio = os.path.dirname(ruta_archivo)
        if directorio:
            os.makedirs(directorio, exist_ok=True)

        AgentExecutor._actualizar_progreso(agente, 70, "Escribiendo .pdf...")

        titulo = None
        if isinstance(contenido, dict):
            titulo = contenido.get("titulo") or contenido.get("title")
            for clave in ("contenido", "texto", "respuesta_limpia", "respuesta"):
                if clave in contenido:
                    contenido = contenido[clave]
                    break

        if isinstance(contenido, str):
            texto = contenido
        elif isinstance(contenido, (dict, list)):
            texto = json.dumps(contenido, indent=2, ensure_ascii=False, default=str)
        else:
            texto = str(contenido)

        if not texto.strip():
            return False, f"File.escribir_pdf: contenido vacío para '{ruta_archivo}'", {
                "error": "empty_content", "archivo": ruta_archivo
            }

        doc = SimpleDocTemplate(
            ruta_archivo,
            leftMargin=2 * cm, rightMargin=2 * cm,
            topMargin=2 * cm, bottomMargin=2 * cm,
        )
        styles = getSampleStyleSheet()
        story = []

        if titulo:
            story.append(Paragraph(_escapar_xml(str(titulo)), styles["Title"]))
            story.append(Spacer(1, 0.5 * cm))

        for par in texto.split("\n"):
            if cancellation_token and cancellation_token.esta_cancelado():
                return False, "Cancelado durante generación de .pdf", {
                    "error": "cancelled", "archivo": ruta_archivo
                }
            if not par.strip():
                story.append(Spacer(1, 0.2 * cm))
                continue
            story.append(Paragraph(_escapar_xml(par), styles["Normal"]))

        if not story:
            story = [Paragraph("(documento vacío)", styles["Normal"])]

        try:
            doc.build(story)
        except Exception as e:
            return False, f"File.escribir_pdf: error al construir '{ruta_archivo}': {e}", {
                "error": "pdf_build_error", "archivo": ruta_archivo, "detalle": str(e)
            }

        if not os.path.exists(ruta_archivo) or os.path.getsize(ruta_archivo) == 0:
            return False, f"File.escribir_pdf: '{ruta_archivo}' no se creó o está vacío", {
                "error": "write_failed", "archivo": ruta_archivo
            }

        tamaño = os.path.getsize(ruta_archivo)
        AgentExecutor._actualizar_progreso(agente, 100, "Archivo .pdf escrito")

        return True, f"Archivo .pdf escrito: {ruta_archivo} ({tamaño} bytes)", {
            "archivo": ruta_archivo,
            "ruta_absoluta": os.path.abspath(ruta_archivo),
            "tamaño": tamaño,
            "bytes_en_disco": tamaño,
            "formato": "pdf",
        }

    @staticmethod
    def _file_copiar(agente: Agente, origen: str, destino: str) -> Tuple[bool, str, Dict]:
        if not os.path.exists(origen):
            AgentExecutor._actualizar_progreso(agente, 100, "Origen no encontrado")
            return False, f"Archivo origen no encontrado: {origen}", {}
        if not os.path.isfile(origen):
            AgentExecutor._actualizar_progreso(agente, 100, "No es un archivo")
            return False, f"No es un archivo: {origen}", {}

        directorio_destino = os.path.dirname(destino)
        if directorio_destino:
            os.makedirs(directorio_destino, exist_ok=True)

        AgentExecutor._actualizar_progreso(agente, 70, "Copiando archivo...")
        shutil.copy2(origen, destino)
        AgentExecutor._actualizar_progreso(agente, 100, "Archivo copiado")

        return True, f"Archivo copiado: {origen} → {destino}", {
            "origen": origen,
            "destino": destino,
            "tamaño": os.path.getsize(destino),
        }

    @staticmethod
    def _file_mover(agente: Agente, origen: str, destino: str) -> Tuple[bool, str, Dict]:
        if not os.path.exists(origen):
            AgentExecutor._actualizar_progreso(agente, 100, "Origen no encontrado")
            return False, f"Archivo origen no encontrado: {origen}", {}
        if not os.path.isfile(origen):
            AgentExecutor._actualizar_progreso(agente, 100, "No es un archivo")
            return False, f"No es un archivo: {origen}", {}

        directorio_destino = os.path.dirname(destino)
        if directorio_destino:
            os.makedirs(directorio_destino, exist_ok=True)

        AgentExecutor._actualizar_progreso(agente, 70, "Moviendo archivo...")
        shutil.move(origen, destino)
        AgentExecutor._actualizar_progreso(agente, 100, "Archivo movido")

        return True, f"Archivo movido: {origen} → {destino}", {
            "origen": origen,
            "destino": destino,
            "tamaño": os.path.getsize(destino) if os.path.exists(destino) else None,
        }

    @staticmethod
    def _file_eliminar(agente: Agente, ruta_archivo: str) -> Tuple[bool, str, Dict]:
        if not os.path.exists(ruta_archivo):
            AgentExecutor._actualizar_progreso(agente, 100, "Archivo no encontrado")
            return False, f"Archivo no encontrado: {ruta_archivo}", {}

        if os.path.isdir(ruta_archivo):
            if ruta_archivo in DANGEROUS_DIRS or os.path.dirname(ruta_archivo) in DANGEROUS_DIRS:
                AgentExecutor._actualizar_progreso(agente, 100, "No se permite eliminar")
                return False, (
                    f"No se permite eliminar directorios del sistema: {ruta_archivo}"
                ), {}
            AgentExecutor._actualizar_progreso(agente, 70, "Eliminando directorio...")
            shutil.rmtree(ruta_archivo)
            mensaje = f"Directorio eliminado: {ruta_archivo}"
        else:
            AgentExecutor._actualizar_progreso(agente, 70, "Eliminando archivo...")
            os.remove(ruta_archivo)
            mensaje = f"Archivo eliminado: {ruta_archivo}"

        AgentExecutor._actualizar_progreso(agente, 100, "Eliminado")
        return True, mensaje, {"archivo": ruta_archivo, "eliminado": True}

    # ============================================================
    # HELPERS DE ESCRITURA
    # ============================================================

    @staticmethod
    def _extraer_contenido_para_escritura(
        contexto: Dict,
        operacion: str,
    ) -> Tuple[Optional[Any], Optional[Tuple[bool, str, Dict]]]:
        """
        Devuelve (contenido, None) si hay contenido válido, o
        (None, (False, mensaje, resultado_error)) si hay que abortar.
        """
        contenido = None

        if "contenido" in contexto:
            contenido = contexto["contenido"]
        elif "resultado" in contexto:
            contenido = contexto["resultado"]
        elif contexto:
            if len(contexto) == 1:
                contenido = AgentExecutor._extraer_contenido_relevante(
                    next(iter(contexto.values()))
                )
            else:
                for clave, valor in contexto.items():
                    if not isinstance(valor, dict):
                        continue
                    candidato = AgentExecutor._extraer_contenido_relevante(valor)
                    if candidato is not valor:
                        contenido = candidato
                        break
                if contenido is None:
                    contenido = AgentExecutor._extraer_contenido_relevante(
                        next(iter(contexto.values()))
                    )

        if contenido is None:
            claves = list(contexto.keys()) if contexto else []
            return None, (False, (
                f"File.{operacion}: no se encontró contenido para escribir. "
                f"Claves del contexto: {claves if claves else '(vacío)'}"
            ), {
                "error": "no_content",
                "operacion": operacion,
                "contexto_claves": claves,
            })

        if isinstance(contenido, str) and not contenido.strip():
            return None, (False, f"File.{operacion}: contenido vacío (string vacío)", {
                "error": "empty_content", "operacion": operacion
            })
        if isinstance(contenido, (dict, list)) and not contenido:
            return None, (False, (
                f"File.{operacion}: contenido vacío ({type(contenido).__name__} vacío)"
            ), {
                "error": "empty_content", "operacion": operacion
            })

        return contenido, None

    @staticmethod
    def _aplicar_modo_salida_file(resultado: Dict, modo: str, contenido_texto: str) -> Dict:
        """Ajusta el resultado según el modo elegido (auto|contenido|texto|json)."""
        modo = (modo or "auto").lower()

        if modo == "auto":
            return resultado

        if modo in ("contenido", "texto"):
            base = dict(resultado)
            base.update({
                "archivo": resultado.get("archivo", ""),
                "tamaño": resultado.get(
                    "tamaño", resultado.get("bytes_en_disco", len(contenido_texto))
                ),
                "bytes_en_disco": resultado.get(
                    "bytes_en_disco", resultado.get("tamaño", len(contenido_texto))
                ),
                "contenido": contenido_texto,
                "total_caracteres": len(contenido_texto),
                "modo_salida": "contenido",
            })
            return base

        if modo == "json":
            json_data = resultado.get("json")
            if json_data is None:
                try:
                    json_data = json.loads(contenido_texto)
                except (json.JSONDecodeError, ValueError):
                    json_data = None
            resultado_json = {
                "archivo": resultado.get("archivo", ""),
                "tamaño": resultado.get(
                    "tamaño", resultado.get("bytes_en_disco", len(contenido_texto))
                ),
                "bytes_en_disco": resultado.get(
                    "bytes_en_disco", resultado.get("tamaño", len(contenido_texto))
                ),
                "json": json_data,
                "modo_salida": "json",
            }
            if json_data is None:
                resultado_json["contenido"] = contenido_texto
            return resultado_json

        return resultado

    # ============================================================
    # 5. LLM (DeepSeek)
    # ============================================================

    @staticmethod
    def _ejecutar_llm(
        agente: Agente,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        if cancellation_token and cancellation_token.esta_cancelado():
            return False, "Cancelado antes de ejecutar", {"error": "cancelled"}

        AgentExecutor._actualizar_progreso(agente, 20, "Preparando prompt LLM...")

        if not agente.prompt_llm:
            AgentExecutor._actualizar_progreso(agente, 100, "Prompt vacío")
            return False, "No hay prompt definido", {}

        try:
            import openai
        except ImportError:
            AgentExecutor._actualizar_progreso(agente, 100, "openai no instalado")
            return False, "openai no instalado. Ejecuta: pip install openai", {}

        variables = AgentExecutor._variables_disponibles(agente, contexto)
        prompt_procesado = AgentExecutor._sustituir_variables(agente.prompt_llm, variables)
        se_sustituyo_algo = prompt_procesado != agente.prompt_llm

        try:
            from core.llm_client import obtener_llm_client_compartido
            _cliente_llm = obtener_llm_client_compartido()
            api_key = _cliente_llm.api_key
            base_url = _cliente_llm.base_url
        except Exception as e:
            AgentExecutor._actualizar_progreso(agente, 100, "Error cargando config LLM")
            return False, (
                f"No se pudo inicializar el cliente LLM: {e}\n"
                f"Prueba: python main.py --check-env"
            ), {"error": "llm_client_init", "detalle": str(e)}

        if not api_key:
            AgentExecutor._actualizar_progreso(agente, 100, "Falta DEEPSEEK_API_KEY")
            return False, (
                "No se encontró DEEPSEEK_API_KEY.\nPrueba: python main.py --check-env"
            ), {"error": "missing_api_key"}

        modelo = getattr(agente, "modelo_llm", "deepseek-v4-flash")
        reasoning_effort = getattr(agente, "reasoning_effort_llm", "low") or "low"
        thinking_enabled = bool(getattr(agente, "thinking_enabled_llm", False))

        MIN_TOKENS_SEGUROS = 4000
        max_tokens_efectivos = int(
            getattr(agente, "max_tokens_llm", MIN_TOKENS_SEGUROS) or MIN_TOKENS_SEGUROS
        )
        if max_tokens_efectivos < MIN_TOKENS_SEGUROS:
            logger.warning(
                f"⚠ [{agente.nombre}] max_tokens={max_tokens_efectivos} bajo, "
                f"subiendo a {MIN_TOKENS_SEGUROS}."
            )
            max_tokens_efectivos = MIN_TOKENS_SEGUROS

        AgentExecutor._actualizar_progreso(agente, 40, f"Consultando {modelo}...")

        if cancellation_token and cancellation_token.esta_cancelado():
            return False, "Cancelado antes de llamar a la API", {"error": "cancelled"}

        try:
            AgentExecutor._get_rate_limiter().wait()

            client = (
                getattr(_cliente_llm, "openai_client", None)
                or getattr(_cliente_llm, "_client", None)
            )
            if client is None:
                client = openai.OpenAI(api_key=api_key, base_url=base_url)

            messages = [
                {
                    "role": "system",
                    "content": (
                        "Eres un asistente útil y preciso. "
                        "Responde directamente con lo que se te pide, sin preámbulos."
                    ),
                },
                {"role": "user", "content": prompt_procesado},
            ]

            if contexto and not se_sustituyo_algo:
                visible = {
                    k: v for k, v in contexto.items()
                    if k not in ("llm_base_url", "llm_api_key", "openai_api_key")
                }
                if visible:
                    messages[1]["content"] += (
                        "\n\nContexto adicional:\n"
                        + json.dumps(visible, indent=2, default=str, ensure_ascii=False)
                    )

            AgentExecutor._actualizar_progreso(agente, 60, "Esperando respuesta...")

            start_time = time.time()
            response = None
            error = None
            completed = threading.Event()

            def hacer_llamada():
                nonlocal response, error
                try:
                    response = client.chat.completions.create(
                        model=modelo,
                        messages=messages,
                        temperature=agente.temperatura_llm,
                        max_tokens=max_tokens_efectivos,
                        reasoning_effort=reasoning_effort,
                        stream=False,
                        extra_body={
                            "thinking": {
                                "type": "enabled" if thinking_enabled else "disabled"
                            }
                        },
                        timeout=60,
                    )
                except Exception as e:
                    error = e
                finally:
                    completed.set()

            threading.Thread(target=hacer_llamada, daemon=True).start()

            timeout = 60
            inicio = time.time()
            while not completed.is_set():
                if cancellation_token and cancellation_token.esta_cancelado():
                    AgentExecutor._actualizar_progreso(agente, 100, "⛔ Cancelado")
                    return False, "Cancelado por usuario", {
                        "error": "cancelled", "modelo": modelo,
                        "tiempo_espera": time.time() - inicio,
                    }
                if time.time() - inicio > timeout + 5:
                    AgentExecutor._actualizar_progreso(agente, 100, "⏱ Timeout")
                    return False, f"Timeout en llamada LLM ({timeout}s)", {
                        "error": "timeout", "modelo": modelo, "tiempo_espera": timeout,
                    }
                completed.wait(0.1)

            if error:
                raise error
            if response is None or not response.choices:
                raise openai.APIError("No se recibió respuesta de la API")

            elapsed_time = time.time() - start_time
            respuesta = response.choices[0].message.content if response.choices else ""

            if not respuesta or not respuesta.strip():
                reasoning = getattr(response.choices[0].message, "reasoning_content", None)
                if reasoning:
                    tokens_info = {}
                    if response.usage:
                        tokens_info = {
                            "prompt": response.usage.prompt_tokens,
                            "completion": response.usage.completion_tokens,
                            "total": response.usage.total_tokens,
                        }
                    return False, (
                        f"⚠ LLM solo devolvió razonamiento (no respuesta final). "
                        f"Aumenta max_tokens (actual: {agente.max_tokens_llm}). "
                        f"Tokens: {tokens_info.get('total', '?')}"
                    ), {
                        "error": "only_reasoning", "modelo": modelo,
                        "razonamiento_preview": reasoning[:200],
                        "tokens_uso": tokens_info,
                    }
                return False, (
                    f"LLM devolvió respuesta vacía (modelo: {modelo})"
                ), {"error": "empty_response", "modelo": modelo}

            prompt_lower = prompt_procesado.lower()
            pide_generar = any(
                kw in prompt_lower for kw in (
                    "genera", "escribe", "redacta", "crea", "mensaje",
                    "texto", "resumen", "explica", "describe", "elabora",
                    "lista", "enumera", "traduce",
                )
            )
            if pide_generar and len(respuesta.strip()) < 3:
                tokens_usados = response.usage.total_tokens if response.usage else None
                return False, (
                    f"⚠ El LLM devolvió una respuesta truncada: '{respuesta}'\n"
                    f"Sugerencia: aumenta max_tokens a 500 o más."
                ), {
                    "error": "truncated_response", "modelo": modelo,
                    "respuesta_truncada": respuesta,
                    "max_tokens": agente.max_tokens_llm,
                    "tokens_usados": tokens_usados,
                }

            palabras_thinking = [
                "We need", "The user", "I need to", "Let me", "First,",
                "Okay,", "Alright,", "Hmm,", "So,", "Now,",
            ]
            respuesta_lower = respuesta.strip()[:50].lower()
            if any(respuesta_lower.startswith(p.lower()) for p in palabras_thinking) \
                    and len(respuesta) > 300:
                logger.warning(
                    f"⚠ LLM parece haber devuelto razonamiento en lugar "
                    f"de respuesta. Modelo: {modelo}, Longitud: {len(respuesta)}"
                )

            AgentExecutor._actualizar_progreso(agente, 90, "Procesando respuesta...")

            respuesta_limpia = _limpiar_fences_markdown(respuesta)
            json_auto = _parsear_json_robusto(respuesta_limpia)

            if json_auto is None and respuesta_limpia.strip().startswith(("{", "[")):
                error_msg = (
                    f"La respuesta del LLM parece ser un JSON truncado "
                    f"({len(respuesta_limpia)} caracteres). "
                    f"Aumenta 'max_tokens_llm' (actual: {agente.max_tokens_llm}) "
                    f"o desactiva 'thinking'."
                )
                logger.error(f"[{agente.nombre}] {error_msg}")
                AgentExecutor._actualizar_progreso(agente, 100, "❌ JSON truncado")
                return False, error_msg, {
                    "error": "truncated_json", "modelo": modelo,
                    "raw_response": respuesta[:500],
                    "total_length": len(respuesta),
                    "max_tokens_actual": agente.max_tokens_llm,
                }

            resultado = {
                "modelo": modelo,
                "prompt": prompt_procesado,
                "respuesta": respuesta,
                "respuesta_limpia": respuesta_limpia,
                "json": json_auto,
                "_av_respuesta": respuesta,
                "_av_respuesta_limpia": respuesta_limpia,
                "_av_json": json_auto,
                "tiempo_respuesta": elapsed_time,
                "tokens_uso": {
                    "prompt": response.usage.prompt_tokens if response.usage else 0,
                    "completion": response.usage.completion_tokens if response.usage else 0,
                    "total": response.usage.total_tokens if response.usage else 0,
                },
            }

            sospechoso, razon = _es_resultado_sospechoso(
                json_auto if json_auto is not None else respuesta_limpia
            )
            if sospechoso and len(respuesta_limpia) < 20:
                logger.warning(f"[{agente.nombre}] Resultado sospechoso: {razon}")

            AgentExecutor._actualizar_progreso(agente, 100, "LLM respondió")
            preview = (
                respuesta_limpia[:200] if len(respuesta_limpia) > 20
                else respuesta[:200]
            )
            return True, f"DeepSeek respondió: {preview}...", resultado

        except openai.APIError as e:
            AgentExecutor._actualizar_progreso(agente, 100, f"Error API")
            return False, f"Error de API de DeepSeek: {e}", {
                "error": "api_error", "detalle": str(e)
            }
        except openai.APIConnectionError as e:
            AgentExecutor._actualizar_progreso(agente, 100, "Error de conexión")
            return False, f"Error de conexión con DeepSeek: {e}", {
                "error": "connection_error", "detalle": str(e)
            }
        except openai.APITimeoutError as e:
            AgentExecutor._actualizar_progreso(agente, 100, "Timeout")
            return False, f"Timeout en DeepSeek: {e}", {
                "error": "timeout", "detalle": str(e)
            }
        except openai.RateLimitError as e:
            AgentExecutor._actualizar_progreso(agente, 100, "Rate limit")
            return False, f"Rate limit excedido en DeepSeek: {e}", {
                "error": "rate_limit", "detalle": str(e)
            }
        except Exception as e:
            AgentExecutor._actualizar_progreso(agente, 100, f"Error: {str(e)[:50]}")
            return False, f"Error en DeepSeek: {e}", {
                "error": "unexpected", "detalle": str(e)
            }

    # ============================================================
    # 6. LOOP
    # ============================================================

    @staticmethod
    def _ejecutar_loop(
        agente: Agente,
        contexto: Dict,
        cancellation_token: Optional[CancellationToken] = None,
    ) -> Tuple[bool, str, Dict]:
        if cancellation_token and cancellation_token.esta_cancelado():
            return False, "Cancelado antes de ejecutar", {"error": "cancelled"}

        AgentExecutor._actualizar_progreso(agente, 10, "Preparando loop...")

        if not agente.fuente_items:
            AgentExecutor._actualizar_progreso(agente, 100, "Fuente de items vacía")
            return False, "Loop: 'fuente_items' no está configurado", {}
        if not agente.codigo_por_item or not agente.codigo_por_item.strip():
            AgentExecutor._actualizar_progreso(agente, 100, "Código por item vacío")
            return False, "Loop: 'codigo_por_item' está vacío", {}

        items = AgentExecutor._resolver_ruta_en_contexto(contexto, agente.fuente_items)
        if not AgentExecutor._es_lista_valida(items):
            AgentExecutor._actualizar_progreso(agente, 100, "Fuente no es una lista")
            return False, (
                f"Loop: '{agente.fuente_items}' no resolvió a lista "
                f"(obtuve: {type(items).__name__})"
            ), {"fuente": agente.fuente_items, "valor": items}

        total_items = len(items)
        if total_items == 0:
            AgentExecutor._actualizar_progreso(agente, 100, "Lista vacía")
            return True, "Loop: La lista está vacía", {
                "total_items": 0, "items_procesados": 0,
                "exitos": 0, "errores": 0, "no_ejecutados": 0,
            }

        if total_items > agente.max_iteraciones:
            AgentExecutor._actualizar_progreso(agente, 100, "Excede máx. iteraciones")
            return False, (
                f"Loop: {total_items} items > máx {agente.max_iteraciones}"
            ), {"total_items": total_items, "max_iteraciones": agente.max_iteraciones}

        resultados = []
        errores = 0
        exitos = 0
        tiempo_inicio = time.time()
        timeout_total = agente.timeout_loop or 300
        timeout_item = agente.timeout_python or 30

        AgentExecutor._actualizar_progreso(agente, 15, f"Procesando {total_items} items...")

        for idx, item in enumerate(items):
            if cancellation_token and cancellation_token.esta_cancelado():
                AgentExecutor._actualizar_progreso(agente, 100, "⛔ Cancelado")
                return False, f"Loop cancelado en item {idx + 1}/{total_items}", {
                    "error": "cancelled",
                    "items_procesados": idx,
                    "total_items": total_items,
                    "errores": errores, "exitos": exitos,
                    "no_ejecutados": total_items - idx,
                    "resultados_parciales": resultados,
                }

            if time.time() - tiempo_inicio > timeout_total:
                AgentExecutor._actualizar_progreso(agente, 100, "Timeout global")
                return False, (
                    f"Loop: Timeout global ({timeout_total}s) tras {idx} items"
                ), {
                    "items_procesados": idx,
                    "total_items": total_items,
                    "errores": errores, "exitos": exitos,
                    "no_ejecutados": total_items - idx,
                    "resultados_parciales": resultados,
                }

            progreso = 15 + int((idx / total_items) * 75)
            AgentExecutor._actualizar_progreso(
                agente, progreso,
                f"Item {idx + 1}/{total_items}: {str(item)[:30]}..."
            )

            ctx_item = dict(contexto)
            ctx_item["item"] = item
            ctx_item["indice"] = idx
            ctx_item["total"] = total_items

            try:
                exito, msg, res = PythonSandbox.ejecutar(
                    agente.codigo_por_item, ctx_item,
                    timeout=timeout_item, cancellation_token=cancellation_token,
                )
                resultados.append({
                    "indice": idx, "item": item,
                    "exito": exito, "mensaje": msg, "resultado": res,
                })
                if exito:
                    exitos += 1
                else:
                    errores += 1
                    if not agente.continuar_en_error:
                        duracion = time.time() - tiempo_inicio
                        no_ejec = total_items - len(resultados)
                        AgentExecutor._actualizar_progreso(
                            agente, 100, f"⛔ Detenido en item {idx + 1}"
                        )
                        return False, (
                            f"Loop detenido en item {idx + 1}: {msg[:200]}"
                        ), {
                            "total_items": total_items,
                            "items_procesados": len(resultados),
                            "exitos": exitos, "errores": errores,
                            "no_ejecutados": no_ejec,
                            "duracion_total": duracion,
                            "detenido_en_indice": idx,
                            "detenido_por_error": True,
                            "continuar_en_error": False,
                            "mensaje_error": msg[:500],
                            "items": resultados,
                        }
            except Exception as e:
                errores += 1
                resultados.append({
                    "indice": idx, "item": item,
                    "exito": False, "mensaje": f"Error crítico: {e}",
                    "resultado": {"error": str(e)},
                })
                if not agente.continuar_en_error:
                    duracion = time.time() - tiempo_inicio
                    no_ejec = total_items - len(resultados)
                    AgentExecutor._actualizar_progreso(
                        agente, 100, f"⛔ Detenido en item {idx + 1}"
                    )
                    return False, (
                        f"Loop detenido en item {idx + 1}: {e}"
                    ), {
                        "total_items": total_items,
                        "items_procesados": len(resultados),
                        "exitos": exitos, "errores": errores,
                        "no_ejecutados": no_ejec,
                        "duracion_total": duracion,
                        "detenido_en_indice": idx,
                        "detenido_por_error": True,
                        "continuar_en_error": False,
                        "mensaje_error": str(e)[:500],
                        "items": resultados,
                    }

        duracion = time.time() - tiempo_inicio
        resultado_final = {
            "total_items": total_items,
            "items_procesados": len(resultados),
            "exitos": exitos, "errores": errores,
            "no_ejecutados": 0,
            "duracion_total": duracion,
            "detenido_por_error": False,
            "continuar_en_error": agente.continuar_en_error,
            "items": resultados,
        }

        mensaje = (
            f"Loop completado: {total_items} items, "
            f"{exitos} éxitos, {errores} errores, "
            f"duración: {duracion:.2f}s"
        )
        if errores > 0 and agente.continuar_en_error:
            mensaje += " (continuar_en_error activo)"

        AgentExecutor._actualizar_progreso(agente, 100, "Loop completado")
        exito_general = errores == 0 or agente.continuar_en_error
        return exito_general, mensaje, resultado_final

    # ============================================================
    # 7. HELPERS DE TEST
    # ============================================================

    @staticmethod
    def probar_loop(agente: Agente, items: List[Any]) -> Tuple[bool, str, Dict]:
        if agente.tipo != TipoAgente.LOOP:
            return False, "El agente no es de tipo LOOP", {}
        es_valido, mensaje = agente.validar_configuracion()
        if not es_valido:
            return False, f"Configuración inválida: {mensaje}", {}

        partes = agente.fuente_items.split(".")
        nombre_dep = partes[0]
        clave_items = ".".join(partes[1:]) if len(partes) > 1 else "items"

        contexto = {nombre_dep: {clave_items: items}, "items_prueba": items}
        return AgentExecutor._ejecutar_loop(agente, contexto)

    @classmethod
    def get_status(cls) -> Dict:
        return {"http_cache": cls.get_http_cache_stats(), "rate_limiter": "active"}


# ============================================================
# LIMPIEZA
# ============================================================

@atexit.register
def _cleanup_executor():
    try:
        AgentExecutor.clear_http_cache()
    except Exception:
        pass